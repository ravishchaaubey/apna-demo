"""number_check.py: compare every number in the storyboard's number-consistency table with the deck.

    python3 tools/number_check.py ../slide_storyboard.md ../Priyanka_Nanda_Academy_Growth_Strategy.pptx

For each table row, numeric tokens in the "Value" column are looked up first on the slides the row
cites (e.g. "(5, A1)"), then anywhere in the deck. Slide text, table cells and native chart data
(categories and series values) count; speaker notes are reported separately and do not count.
"""
import re
import sys
from pptx import Presentation

NUM = re.compile(r"\d+(?:,\d{3})*(?:\.\d+)?")


def norm(t):
    return t.replace(",", "")


def fmt_val(v):
    if v is None:
        return []
    out = {f"{abs(v):g}", f"{abs(v):.1f}", f"{abs(v):.2f}"}
    return list(out)


def slide_text(slide):
    parts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        if getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            for plot in ch.plots:
                parts.extend(str(c) for c in plot.categories)
                for s in plot.series:
                    parts.append(s.name)
                    for v in s.values:
                        parts.extend(fmt_val(v))
    return "\n".join(parts)


def tokens(text):
    return {norm(t) for t in NUM.findall(text)}


def main(md_path, pptx_path):
    md = open(md_path, encoding="utf-8").read()
    table = md[md.index("## Number-consistency table"):]
    rows = [l for l in table.splitlines() if re.match(r"^\| \d+ \|", l)]
    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    texts = [slide_text(s) for s in slides]
    notes = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else "" for s in slides]
    # Map storyboard slide labels to physical slide indices
    label_map = {}
    for i, t in enumerate(texts):
        m = re.search(r"APPENDIX (A\d)", t)
        if m:
            label_map.setdefault(m.group(1), []).append(i)
    for k in range(1, 25):
        label_map[str(k)] = [k - 1]
    all_tokens = set().union(*(tokens(t) for t in texts))
    note_tokens = set().union(*(tokens(t) for t in notes))
    results = []
    for line in rows:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        num, head, value = cells[0], cells[1], cells[2]
        m = re.findall(r"\(([^()]*)\)", head)
        refs = []
        if m:
            refs = [r.strip() for r in re.split(r",", m[-1]) if re.match(r"^\s*(A?\d+)\s*$", r)]
        idx = sorted({i for r in refs for i in label_map.get(r, [])})
        cited = set().union(*(tokens(texts[i]) for i in idx)) if idx else set()
        clean = value.replace("**", "")
        vt = [(norm(m.group(0)), clean[m.end():m.end() + 1]) for m in NUM.finditer(clean)]
        on_cited, elsewhere, notes_only, missing, rounded = [], [], [], [], []
        for t, suffix in vt:
            alts = []
            if suffix == "k" and "." in t:  # "₹19.8k" shown as "19,800"
                alts.append(str(int(round(float(t) * 1000))))
            if "." in t and len(t.split(".")[1]) == 2:  # storyboard slide text rounds 11.51 -> 11.5
                alts.append(f"{float(t):.1f}")
            if "." in t:  # "≈7 fewer" for 7.1
                alts.append(str(int(round(float(t)))))
            if t in cited:
                on_cited.append(t)
            elif any(a in cited for a in alts):
                rounded.append(f"{t}~{next(a for a in alts if a in cited)}")
            elif t in all_tokens:
                elsewhere.append(t)
            elif t in note_tokens:
                notes_only.append(t)
            else:
                missing.append(t)
        if notes_only or missing:
            status = "CHECK"
        elif elsewhere:
            status = "MATCH (other slide)"
        elif rounded:
            status = "MATCH (format/rounding)"
        else:
            status = "MATCH"
        results.append((num, head, refs, status, elsewhere, notes_only, missing, rounded))
    for r in results:
        num, head, refs, status, elsewhere, notes_only, missing, rounded = r
        extra = []
        if rounded:
            extra.append("shown as: " + ", ".join(rounded))
        if elsewhere:
            extra.append("on other slides: " + ", ".join(elsewhere))
        if notes_only:
            extra.append("notes only: " + ", ".join(notes_only))
        if missing:
            extra.append("NOT FOUND: " + ", ".join(missing))
        print(f"{num:>3} {status:<20} slides {','.join(refs) or '-':<10} {head[:60]:<60} {'; '.join(extra)}")
    n_ok = sum(1 for r in results if r[3] == "MATCH")
    count = lambda st: sum(1 for r in results if r[3] == st)
    print(f"\n{len(results)} rows: {n_ok} exact on the cited slides; {count('MATCH (format/rounding)')} same value in "
          f"another format or the storyboard's rounding; {count('MATCH (other slide)')} with a token on a different slide; "
          f"{count('CHECK')} need review")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
