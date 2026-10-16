"""Export assumption provenance tables (markdown) from the recalculated workbooks.

Reads the Assumptions sheets of all three workbooks plus the BM_Unit_Economics input block and writes
_provenance_full.md (every ID, for assumptions_and_methodology.md) and _register_rows.md (consolidated,
seasonality/budget rows collapsed, for assumptions_register.md). Values are read after LibreOffice recalc.
"""
import os
import re

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ECON = os.path.dirname(HERE)
BOOKS = [("market_sizing_model.xlsx", "MS"), ("funnel_economics.xlsx", "FE"), ("business_forecast.xlsx", "BF")]


def fmt(v, numfmt):
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if "%" in (numfmt or ""):
        x = v * 100
        if abs(x - round(x)) < 1e-9:
            return f"{x:.0f}%"
        if abs(x - round(x, 1)) < 1e-9:
            return f"{x:.1f}%"
        return f"{x:.2f}%"
    if "₹" in (numfmt or ""):
        return f"₹{v:,.0f}"
    if isinstance(v, float) and not v.is_integer():
        return f"{v:,.2f}"
    return f"{v:,.0f}"


def collect():
    rows = {}
    for fn, tag in BOOKS:
        wb = openpyxl.load_workbook(os.path.join(ECON, fn), data_only=True)
        ws = wb["Assumptions"]
        for r in range(7, ws.max_row + 1):
            aid = ws.cell(row=r, column=1).value
            if not aid or not isinstance(aid, str) or ws.cell(row=r, column=2).value is None:
                continue
            nf = ws.cell(row=r, column=4).number_format
            vals = tuple(fmt(ws.cell(row=r, column=c).value, nf) for c in (3, 4, 5))
            entry = rows.setdefault(aid, dict(desc=ws.cell(row=r, column=2).value, vals=vals, unit=ws.cell(row=r, column=6).value,
                                              src=ws.cell(row=r, column=7).value, tag=ws.cell(row=r, column=8).value, books=[]))
            if tag not in entry["books"]:
                entry["books"].append(tag)
        if fn == "business_forecast.xlsx":
            wbm = wb["BM_Unit_Economics"]
            for r in range(5, 30):
                aid = wbm.cell(row=r, column=1).value
                if aid and str(aid).startswith("A-ECO-3"):
                    v = wbm.cell(row=r, column=3).value
                    unit = wbm.cell(row=r, column=4).value
                    fv = fmt(v, "0%" if unit == "%" else "")
                    rows[aid] = dict(desc=wbm.cell(row=r, column=2).value, vals=("", fv, ""), unit=unit,
                                     src=wbm.cell(row=r, column=5).value, tag="[HYP]", books=["BF (BM_Unit_Economics, Base only)"])
    return rows


def esc(s):
    return str(s or "").replace("|", "/").replace("\n", " ")


def main():
    rows = collect()
    order = sorted(rows, key=lambda k: (k.split("-")[0] != "CA", re.sub(r"(\d+)", lambda m: m.group(1).zfill(4), k)))
    full = ["| ID | Assumption | Low / Cons. | Base | High / Aggr. | Unit | Source / rationale | Tag | Workbook(s) |", "|---|---|---|---|---|---|---|---|---|"]
    for k in order:
        e = rows[k]
        full.append(f"| {k} | {esc(e['desc'])} | {e['vals'][0]} | {e['vals'][1]} | {e['vals'][2]} | {esc(e['unit'])} | {esc(e['src'])} | {esc(e['tag'])} | {', '.join(e['books'])} |")
    open(os.path.join(HERE, "_provenance_full.md"), "w").write("\n".join(full) + "\n")

    # consolidated: collapse series
    reg = ["| ID | Assumption | Low / Cons. | Base | High / Aggr. | Unit | Source / rationale | Tag |", "|---|---|---|---|---|---|---|---|"]
    seen_series = set()
    for k in order:
        e = rows[k]
        m = re.match(r"^(A-GRO-40|A-ECO-60|A-GRO-BUD)-(\d\d)$", k)
        if m:
            if m.group(1) in seen_series:
                continue
            seen_series.add(m.group(1))
            series = [rows[x] for x in order if x.startswith(m.group(1) + "-")]
            if m.group(1) == "A-GRO-BUD":
                vals = tuple("/".join(s["vals"][j].replace("₹", "").replace(",000", "k") for s in series) for j in range(3))
                reg.append(f"| A-GRO-BUD-01..12 | Paid-media budget by month, Nov-26..Oct-27 | {vals[0]} | {vals[1]} | {vals[2]} | ₹ | growth_marketing_strategy.md §6.4/§10 | [EST] |")
            else:
                vals = "/".join(s["vals"][1] for s in series)
                nm = "Career-enquiry seasonality Jan..Dec" if m.group(1) == "A-GRO-40" else "Self-makeup seasonality Jan..Dec"
                reg.append(f"| {m.group(1)}-01..12 | {nm} | — | {vals} | — | index | {esc(series[0]['src'])} | {esc(series[0]['tag'])} |")
            continue
        mc = re.match(r"^(A-ECO-164|A-ECO-165)-(\d)$", k)
        if mc:
            if mc.group(1) in seen_series:
                continue
            seen_series.add(mc.group(1))
            series = [rows[x] for x in order if re.match(rf"^{mc.group(1)}-\d$", x)]
            vals = " / ".join(x["vals"][1] for x in series)
            nm = ("L3 local cohort start times, 6-weekly intake calendar (model months from 1 Nov 2026; 2.10 = 4 Jan 2027)" if mc.group(1) == "A-ECO-164"
                  else "Alternative 8-weekly intake calendar: L3 cohort start times (model months)")
            reg.append(f"| {mc.group(1)}-1..{len(series)} | {nm} | — | {vals} | — | model months | {esc(series[0]['src'])} | {esc(series[0]['tag'])} |")
            continue
        if k.startswith("SOM-"):
            continue
        reg.append(f"| {k} | {esc(e['desc'])} | {e['vals'][0]} | {e['vals'][1]} | {e['vals'][2]} | {esc(e['unit'])} | {esc(e['src'])} | {esc(e['tag'])} |")
    reg.append("| SOM-* | Year-3 revenue / learners by layer pasted from business_forecast.xlsx into market_sizing_model.xlsx | — | — | — | ₹ lakh / learners | Model output (run_all.py) | [EST] |")
    open(os.path.join(HERE, "_register_rows.md"), "w").write("\n".join(reg) + "\n")
    print(len(rows), "ids;", len(reg) - 2, "register rows")


if __name__ == "__main__":
    main()
