"""Shared helpers for the PNA economics workbooks (openpyxl).

Conventions used in every workbook
- Blue font = hard-coded input (Assumptions sheet only, plus a few labelled scenario switches).
- Black font = formula. Every calculation cell is an Excel formula that refers back to the
  Assumptions sheet, so the models are auditable and recalculate in Excel / LibreOffice.
- Assumption rows carry: ID | description | case values | unit | source / rationale | tag.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

NAVY = "1F3864"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(color="FFFFFF", bold=True)
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
SEC_FONT = Font(bold=True, color=NAVY)
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(italic=True, color="595959", size=9)
INPUT_FONT = Font(color="0000CC")
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

FMT_INT = '#,##0;[Red]-#,##0'
FMT_1DP = '#,##0.0;[Red]-#,##0.0'
FMT_2DP = '#,##0.00;[Red]-#,##0.00'
FMT_PCT = '0.0%;[Red]-0.0%'
FMT_PCT2 = '0.00%;[Red]-0.00%'
FMT_RS = '"₹"#,##0;[Red]-"₹"#,##0'
FMT_X = '0.00"x"'
FMT_GEN = 'General'


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if sub:
        ws["A2"] = sub
        ws["A2"].font = NOTE_FONT


def header(ws, row, labels, col=1, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(col + i)].width = w


def section(ws, row, text, ncols=8, col=1):
    for i in range(ncols):
        ws.cell(row=row, column=col + i).fill = SEC_FILL
    c = ws.cell(row=row, column=col, value=text)
    c.font = SEC_FONT


def note(ws, row, text, col=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = NOTE_FONT
    return c


def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


class Assumptions:
    """Writes an Assumptions sheet and resolves IDs to cell references.

    Columns: A ID | B Description | C case1 | D case2 | E case3 | F Unit | G Source/rationale | H Tag | I Selected
    The Selected column = INDEX(C:E, scenario index) driven by the selector in C3.
    """

    def __init__(self, ws, case_labels=("Low", "Base", "High"), sheet_name="Assumptions",
                 selector_default="Base", start_row=8, intro=None):
        self.ws = ws
        self.sheet = sheet_name
        self.rows = {}
        self.meta = {}
        self.case_labels = case_labels
        ws.title = sheet_name
        title(ws, "Assumptions register (inputs)",
              intro or "Blue = input. Every model sheet references these cells. Tags: [FACT] sourced; [CLAIM] first-party/unverified; [EST] our estimate with arithmetic; [HYP] hypothesis to validate; [FOUNDER] founder-stated.")
        ws["A3"] = "Scenario selector"
        ws["A3"].font = BOLD
        ws["C3"] = selector_default
        ws["C3"].font = Font(color="C00000", bold=True)
        ws["C3"].fill = KEY_FILL
        ws["D3"] = "<- choose " + " / ".join(case_labels) + " (drives column I 'Selected' and any sheet that uses it)"
        ws["D3"].font = NOTE_FONT
        for i, lab in enumerate(case_labels):
            ws.cell(row=4, column=11 + i, value=lab).font = NOTE_FONT  # K4:M4 list for MATCH
        ws["A4"] = "Scenario index"
        ws["C4"] = f'=MATCH($C$3,$K$4:$M$4,0)'
        ws["C4"].font = BOLD
        dv = DataValidation(type="list", formula1='"' + ",".join(case_labels) + '"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add("C3")
        header(ws, 6, ["ID", "Description", case_labels[0], case_labels[1], case_labels[2],
                       "Unit", "Source / rationale", "Tag", "Selected"],
               widths=[14, 52, 13, 13, 13, 12, 70, 10, 13])
        ws.freeze_panes = "C7"
        self.row = start_row - 1

    def sec(self, text):
        self.row += 1
        section(self.ws, self.row, text, ncols=9)

    def add(self, aid, desc, v1, v2, v3, unit, source, tag, fmt=FMT_GEN):
        self.row += 1
        r = self.row
        ws = self.ws
        ws.cell(row=r, column=1, value=aid).font = BOLD
        ws.cell(row=r, column=2, value=desc).alignment = WRAP
        for j, v in enumerate((v1, v2, v3)):
            c = ws.cell(row=r, column=3 + j, value=v)
            if not (isinstance(v, str) and v.startswith("=")):
                c.font = INPUT_FONT
            c.number_format = fmt
        ws.cell(row=r, column=6, value=unit)
        ws.cell(row=r, column=7, value=source).alignment = WRAP
        ws.cell(row=r, column=8, value=tag)
        c = ws.cell(row=r, column=9, value=f"=INDEX(C{r}:E{r},$C$4)")
        c.number_format = fmt
        c.fill = KEY_FILL
        self.rows[aid] = r
        self.meta[aid] = dict(desc=desc, vals=(v1, v2, v3), unit=unit, source=source, tag=tag)
        return r

    # ---- references -------------------------------------------------------------------------
    def r(self, aid, case="sel"):
        """Absolute reference to an assumption cell. case: 'sel', 1/2/3, 'low'/'base'/'high'."""
        row = self.rows[aid]
        col = {"sel": "I", 1: "C", 2: "D", 3: "E", "low": "C", "base": "D", "high": "E"}[case]
        return f"{self.sheet}!${col}${row}"

    def pick(self, aid, idx_ref):
        """INDEX into the three case columns using a scenario index held in another cell."""
        row = self.rows[aid]
        return f"INDEX({self.sheet}!$C${row}:$E${row},{idx_ref})"

    def case_col(self, j):
        return "CDE"[j]


def add_list_validation(ws, cell, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(cell)


def fmt_row(ws, row, c1, c2, fmt):
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).number_format = fmt
