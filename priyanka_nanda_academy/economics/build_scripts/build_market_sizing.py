"""Build economics/market_sizing_model.xlsx  (TAM / SAM / SOM with tricity-vs-India reconciliation).

All calculation cells are formulas referencing the Assumptions sheet. Low / Base / High are computed
side by side (columns C/D/E); the Sizing_Tree also shows the case chosen in the Assumptions selector.
SOM inputs are Year-3 outputs of business_forecast.xlsx (read from _forecast_outputs.json after that
workbook is recalculated) and are pasted as labelled values.
"""
import json
import os
import sys

import openpyxl
from openpyxl.styles import Font, Alignment

sys.path.insert(0, os.path.dirname(__file__))
from xlhelpers import *  # noqa

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "economics", "market_sizing_model.xlsx")
FC_JSON = os.path.join(os.path.dirname(__file__), "_forecast_outputs.json")

CASES = ("Low", "Base", "High")
COLS = "CDE"


def build():
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    A = Assumptions(wb.create_sheet(), case_labels=CASES, selector_default="Base")

    # ------------------------------------------------------------------ ASSUMPTIONS
    A.sec("1. India TAM inputs (from research/india_market_analysis.md §8-9; values unchanged)")
    A.add("A-IND-01", "B&WSSC-certified candidates per year (all beauty & wellness)", 130000, 170000, 210000, "learners/yr",
          "10.25 lakh/8 yrs=1.28; 13 lakh/9=1.44; 20.7 lakh/10-11=1.88-2.07 (S-IND-02,-04,-05)", "[EST]", FMT_INT)
    A.add("A-IND-02", "Organised share of vocational training (total = certified / share)", 0.40, 0.40, 0.40, "%",
          "KPMG: unorganised ≈60% of vocational education (S-IND-01 p.67)", "[EST]", FMT_PCT)
    A.add("A-IND-03", "Paid (self-funded) share of trainees", 0.64, 0.72, 0.80, "%",
          "60% unorganised (self-paid) + 40% organised x 10/30/50% fee-paying (S-IND-01, S-IND-04)", "[EST]", FMT_PCT)
    A.add("A-IND-04", "Makeup-led share of paid learners", 0.15, 0.25, 0.35, "%",
          "Most uncertain input; 2 of top-5 aspired roles are makeup (S-IND-01 p.53); Meribindiya catalogue ~half makeup (S-IND-37)", "[HYP]", FMT_PCT)
    A.add("A-IND-05a", "Tier T1 entry/short share of makeup-led learners", 0.55, 0.40, 0.25, "%", "india_market_analysis §8.1a", "[HYP]", FMT_PCT)
    A.add("A-IND-05b", "Tier T1 representative fee", 15000, 20000, 25000, "₹", "Online micro 150-6,000; offline basics 10-35k (S-IND-27,-35,-54)", "[EST]", FMT_RS)
    A.add("A-IND-05c", "Tier T2 core professional share", 0.30, 0.35, 0.35, "%", "§8.1a", "[HYP]", FMT_PCT)
    A.add("A-IND-05d", "Tier T2 representative fee", 45000, 55000, 65000, "₹", "YMCA 54,355; Orane online 45,000; VLCC 78,000 (S-IND-25,-28,-33)", "[EST]", FMT_RS)
    A.add("A-IND-05e", "Tier T3 advanced/bridal share", 0.13, 0.20, 0.30, "%", "§8.1a", "[HYP]", FMT_PCT)
    A.add("A-IND-05f", "Tier T3 representative fee", 100000, 130000, 150000, "₹", "VLCC 1.66 lakh; Lakmé 1.2-1.8 lakh (S-IND-32,-34)", "[EST]", FMT_RS)
    A.add("A-IND-05g", "Tier T4 premium share", 0.02, 0.05, 0.10, "%", "§8.1a", "[HYP]", FMT_PCT)
    A.add("A-IND-05h", "Tier T4 representative fee", 250000, 300000, 350000, "₹", "Pearl 4.45 lakh incl GST; DBA 2.25-5.95 lakh (S-IND-23, S-NC-26)", "[EST]", FMT_RS)
    A.add("A-IND-06", "Fee for non-makeup beauty courses", 10000, 13000, 18000, "₹", "Entry beautician/hair/nail courses (S-IND-31,-38)", "[HYP]", FMT_RS)
    A.add("A-IND-07", "Branded chain/franchise centres offering paid beauty training", 300, 450, 600, "centres", "Orane 115+, Lakmé 140+, VLCC 100-200 (claims; S-IND-30,-58,-59; S-NC-01)", "[HYP]", FMT_INT)
    A.add("A-IND-08", "Paid learners per chain centre per year (all courses)", 80, 110, 125, "learners", "Lakmé-Aptech plan 50,000/5 yrs/80 = 125 (S-IND-59)", "[HYP]", FMT_INT)
    A.add("A-IND-09", "Independent paid learners as multiple of chain learners", 1.5, 2.0, 3.0, "x", "KPMG 60:40 unorganised:organised (S-IND-01)", "[HYP]", FMT_1DP)
    A.add("A-IND-10a", "Online / short-course layer: learners", 50000, 150000, 400000, "learners/yr", "Airblack ~35k students in 3+ yrs; Orane micro-courses (S-IND-27,-39,-40)", "[HYP]", FMT_INT)
    A.add("A-IND-10b", "Online / short-course layer: average fee", 600, 1500, 3000, "₹", "Orane ₹590-5,990; Airblack from ₹150", "[HYP]", FMT_RS)
    A.add("A-IND-16", "FX for USD display only", 88, 88, 88, "₹/USD", "Illustrative", "[HYP]", FMT_INT)
    A.add("A-IND-17", "Regional proxy: Chd+PB+HR+HP+J&K share of India incremental B&W workforce need", 0.0765, 0.0765, 0.0765, "%", "KPMG Table 5 (S-IND-01 pp.46-47)", "[FACT]", FMT_PCT2)
    A.add("X-KPMG-NEED", "KPMG incremental vocational training need per year (5.44M / 8)", 679624, 679624, 679624, "learners/yr", "S-IND-01 Table 4 p.46", "[FACT]", FMT_INT)
    A.add("X-BS-SECTOR", "Context only: Beauty & Salon sector size 2025P (NOT the education market)", 76141, 76141, 76141, "₹ crore", "S-IND-01 Table 1 p.15", "[FACT]", FMT_INT)

    A.sec("2. Demography and intensity (research/chandigarh_market_analysis.md; KPMG state tables)")
    A.add("A-CHD-02", "Women aged 17-35, tricity district catchment (Chandigarh UT + SAS Nagar + Panchkula), 2026", 420, 471, 520, "000",
          "Cohort rebuild from Technical Group projections (S-CHD-01, -03, -04); A-CHD-01/02", "[EST]", FMT_INT)
    A.add("X-W-IND", "Women aged 17-35, India, 2026", 222944, 222944, 222944, "000", "S-CHD-01 Table 20 cohort rebuild (A-CHD-08)", "[EST]", FMT_INT)
    A.add("X-W-ROR", "Women 17-35, rest of Punjab+Haryana+HP+J&K (excl. tricity)", 12376, 12376, 12376, "000", "chandigarh_market_analysis §8.3 (12,662 - 286 in tricity)", "[EST]", FMT_INT)
    A.add("X-W-CHD", "Women 17-35: Chandigarh UT", 185, 185, 185, "000", "§2 table", "[EST]", FMT_INT)
    A.add("X-W-MOH", "Women 17-35: SAS Nagar (Mohali) district", 183, 183, 183, "000", "§2 table", "[EST]", FMT_INT)
    A.add("X-W-PKL", "Women 17-35: Panchkula district", 103, 103, 103, "000", "§2 table", "[EST]", FMT_INT)
    A.add("X-W-PB", "Women 17-35: Punjab excl. Mohali", 4275, 4275, 4275, "000", "4,458 - 183", "[EST]", FMT_INT)
    A.add("X-W-HR", "Women 17-35: Haryana excl. Panchkula", 4627, 4627, 4627, "000", "4,730 - 103", "[EST]", FMT_INT)
    A.add("X-W-HP", "Women 17-35: Himachal Pradesh", 1097, 1097, 1097, "000", "§2 table", "[EST]", FMT_INT)
    A.add("X-W-JK", "Women 17-35: J&K", 2377, 2377, 2377, "000", "§2 table", "[EST]", FMT_INT)
    A.add("X-BW-IND", "B&W share of state workforce: India", 0.0241, 0.0241, 0.0241, "%", "S-IND-01 pp.34-35", "[FACT]", FMT_PCT2)
    A.add("X-BW-CHD", "B&W share of state workforce: Chandigarh", 0.0616, 0.0616, 0.0616, "%", "S-IND-01 (highest in India)", "[FACT]", FMT_PCT2)
    A.add("X-BW-PB", "B&W share of state workforce: Punjab", 0.0364, 0.0364, 0.0364, "%", "S-IND-01", "[FACT]", FMT_PCT2)
    A.add("X-BW-HR", "B&W share of state workforce: Haryana", 0.0344, 0.0344, 0.0344, "%", "S-IND-01", "[FACT]", FMT_PCT2)
    A.add("X-BW-HP", "B&W share of state workforce: HP", 0.0202, 0.0202, 0.0202, "%", "S-IND-01", "[FACT]", FMT_PCT2)
    A.add("X-BW-JK", "B&W share of state workforce: J&K", 0.0207, 0.0207, 0.0207, "%", "S-IND-01", "[FACT]", FMT_PCT2)
    A.add("A-CHD-06", "Tricity affluence weight (income-based; cross-check for intensity)", 1.68, 1.85, 2.00, "x India",
          "RBI PCNSDP: Chandigarh 2.28x, Haryana 1.69x, Mohali 1.5x assumed (S-CHD-06)", "[EST]", FMT_2DP)
    A.add("A-ECO-01", "Tricity demand-intensity index used in top-down share (Low = income weight low; Base = KPMG B&W workforce index; High = Chandigarh income multiple)",
          1.68, "=Tricity_Recon!D12", 2.28, "x India", "Base computed on Tricity_Recon from KPMG state B&W shares weighted by women 17-35", "[EST]", FMT_2DP)
    A.add("A-ECO-02", "Urban/rural ratio of paid makeup-academy participation (k)", 3, 5, 8, "x",
          "Makeup academies are urban (S-LC §3 all in cities); rural trainees mostly scheme-funded (S-IND-01 p.52, §7.1)", "[HYP]", FMT_1DP)
    A.add("A-ECO-03", "Urban share of tricity catchment population", 0.68, 0.72, 0.76, "%",
          "Chandigarh ~97% urban (Census 2011, S-CHD-02); Panchkula 55.8% (S-CHD-04); Mohali ~55% assumed; women-weighted", "[EST]", FMT_PCT)
    A.add("A-ECO-04", "Urban share of India population 2026", 0.37, 0.36, 0.35, "%", "Census 2011 31.1% + trend (general knowledge, not re-sourced)", "[EST]", FMT_PCT)
    A.add("A-CHD-04", "Outstation (non-tricity) share of tricity students", 0.15, 0.25, 0.30, "%", "No student-origin data; A-CUS-28 15-35% of enquiries", "[HYP]", FMT_PCT)
    A.add("A-ECO-05", "Rest-of-region demand-intensity index (PB/HR/HP/J&K)", 1.10, "=SAM_North!D9", 1.45, "x India",
          "Base computed from KPMG B&W shares weighted by women; urbanisation ≈ India so no urban factor", "[EST]", FMT_2DP)

    A.add("A-CHD-03", "Annual participation of tricity resident women 17-35 in paid makeup/beauty-artistry course", 0.0012, 0.0020, 0.0030, "%",
          "Calibrated to bridal-MUA listing pipeline (A-CHD-07); unverified", "[HYP]", FMT_PCT2)
    A.add("A-CHD-05", "Average realised fee per tricity learner (blended across course lengths)", 30000, 45000, 60000, "₹",
          "UrbanPro tutor ₹15-36k per 2-3 months; local diploma ₹35-75k (S-CHD-15, S-LC §6)", "[EST]", FMT_RS)

    A.sec("3. Reconciliation and supply-side inputs (Economics agent)")
    A.add("A-ECO-06", "Makeup-led share of CA-05 'makeup/beauty-artistry' learners", 0.60, 0.70, 0.80, "%",
          "CA-05 scope includes hair/skin/nail artistry; tricity providers are makeup-heavy (S-LC §3)", "[HYP]", FMT_PCT)
    A.add("A-ECO-07", "Chain/brand centres in tricity offering makeup courses", 7, 7, 7, "centres",
          "Orane 34A, Lakmé 17E + 34A, VLCC 34A, Knuckout Chd + Zirakpur, Headmasters Mohali (S-LC-08..24, -38)", "[FACT] count of named centres", FMT_INT)
    A.add("A-ECO-08", "Makeup-led learners per chain centre per year", 30, 45, 60, "learners",
          "A-IND-08 (80-125 all courses) x ~40-50% makeup share", "[HYP]", FMT_INT)
    A.add("A-ECO-09", "Independent makeup institutes in tricity", 5, 5, 5, "institutes", "Zuri, Cassia, Fashion Make Fashion, Cleopatra, SS Makeup (S-LC §3)", "[FACT] count", FMT_INT)
    A.add("A-ECO-10", "Makeup-led learners per independent institute per year", 25, 40, 60, "learners",
          "Zuri claims 5,000+ students since 2015 (~450/yr all courses, [CLAIM] S-LC-13) - discounted heavily", "[HYP]", FMT_INT)
    A.add("A-ECO-11", "Founder-led MUA academies in tricity (incl. PNA)", 8, 8, 8, "academies",
          "PNA, Nav Brar, Tanvi KG, Abhishek, Niti Luthra, Parthi Chopra, Meraki, SK Studio (S-LC §3)", "[EST] count", FMT_INT)
    A.add("A-ECO-12", "Learners per founder-led academy per year", 15, 25, 40, "learners", "PNA founder-stated <60 incl. short courses (CA-01)", "[HYP]", FMT_INT)
    A.add("A-ECO-13", "UrbanPro makeup-class listings, Chandigarh", 111, 111, 111, "listings", "S-CHD-15 (9 Oct 2026)", "[FACT]", FMT_INT)
    A.add("A-ECO-14", "Active share of tutor listings x learners per active tutor per year", 0.6, 0.9, 1.2, "learners/listing",
          "30% active x 2/3/4 learners", "[HYP]", FMT_1DP)

    A.sec("4. SAM inputs (Economics agent)")
    A.add("A-ECO-15", "Share of rest-of-region makeup-led learners reachable by an offline Chandigarh hub", 0.10, 0.15, 0.20, "%",
          "Travel matrix: Ambala/Patiala/Solan/Ludhiana/Karnal ≤2.5 h; beyond needs PG (S-CHD §4); Punjab cities chain-served (S-LC §7)", "[HYP]", FMT_PCT)
    A.add("A-ECO-16", "Hybrid-willing share of non-reachable regional makeup-led learners", 0.05, 0.08, 0.12, "%",
          "KPMG survey: 47% prefer blended, little appetite for fully online (S-IND-01 p.56)", "[HYP]", FMT_PCT)
    A.add("A-ECO-17", "Hybrid programme realised fee (quoted, incl. GST)", 40000, 50000, 60000, "₹", "Orane online diploma ₹45-80k (S-IND-28/29); A-GRO parity with L3 later", "[HYP]", FMT_RS)
    A.add("A-ECO-18", "Working makeup artists, tricity (listed + unlisted + salon staff)", 800, 1000, 1300, "artists",
          "WedMeGood 600+ (text) to ~940 (pages x 20) bridal MUAs (S-CHD-12)", "[EST]", FMT_INT)
    A.add("A-ECO-19", "Working makeup artists within reach (rest of region)", 2000, 3000, 4000, "artists",
          "WedMeGood pages: Ludhiana ≤440, Amritsar ≤340, Dehradun ≤440; plus Jalandhar, Patiala, Ambala, Shimla, salon staff (S-CHD §5)", "[EST]", FMT_INT)
    A.add("A-ECO-20", "Share of working artists buying paid CPD/masterclass per year", 0.20, 0.30, 0.40, "%", "A-CUS-31 repeat 25-40%/yr (Airblack claim)", "[HYP]", FMT_PCT)
    A.add("A-ECO-21", "CPD/masterclass spend per participating artist per year (incl. GST)", 8000, 10000, 12000, "₹", "Masterclass anchors ₹4,999-19,999, avg ₹10k (A-GRO-07; S-CUS-34/54)", "[HYP]", FMT_RS)
    A.add("A-ECO-22", "North-India-reachable share of India online/short-course layer (Hindi/Punjabi content)", 0.25, 0.35, 0.45, "%", "Hindi-belt + Punjab share of online buyers; no data", "[HYP]", FMT_PCT)
    A.add("A-ECO-23", "GST rate on PNA training (private, non-recognised course; SAC 9992 commercial training)", 0.18, 0.18, 0.18, "%",
          "GST Council education flyer: private coaching / unrecognised institutions not exempt (gstcouncil.gov.in/sites/default/files/e-version-gst-flyers/GST_Education_Services.pdf, accessed 2026-10-16); PNA-specific status to confirm with CA", "[FACT] rule; [HYP] for PNA", FMT_PCT)

    # ---- SOM inputs from forecast
    fc = json.load(open(FC_JSON)) if os.path.exists(FC_JSON) else None
    A.sec("5. SOM inputs: Year-3 revenue by layer from business_forecast.xlsx '3yr_Scenarios' (₹ lakh, ex-GST). Low=Conservative, Base=Base, High=Aggressive. Pasted values; rebuild via run_all.py")
    layers = [("SOM-TRI", "Y3 tricity-origin flagship revenue (L1 taster, L1-S5 self-makeup, L2, L3 tricity share)", "tri"),
              ("SOM-REG", "Y3 regional destination revenue (L3 outstation share + satellite if gated in)", "reg"),
              ("SOM-HYB", "Y3 hybrid programme revenue", "hyb"),
              ("SOM-CPD", "Y3 working-artist CPD revenue (L4 masterclasses + L5 Pro Circle)", "cpd"),
              ("SOM-ONL", "Y3 online short-course revenue", "onl"),
              ("SOM-ADJ", "Y3 adjacent revenue outside the training market (kit margin, services desk, partnerships)", "adj")]
    for sid, desc, key in layers:
        vals = [fc[s]["y3_" + key] if fc else 0 for s in ("Conservative", "Base", "Aggressive")]
        A.add(sid, desc, *vals, "₹ lakh ex-GST", "business_forecast.xlsx 3yr_Scenarios (linked by run_all.py)", "[EST] model output", FMT_1DP)
    for sid, desc, key in [("SOM-LRN-TRI", "Y3 tricity-origin career learners (L2+L3 new students)", "lrn_tri"),
                           ("SOM-LRN-REG", "Y3 out-of-town career learners (offline + satellite)", "lrn_reg"),
                           ("SOM-LRN-HYB", "Y3 incremental destination/hybrid learners", "lrn_hyb")]:
        vals = [fc[s]["y3_" + key] if fc else 0 for s in ("Conservative", "Base", "Aggressive")]
        A.add(sid, desc, *vals, "learners", "business_forecast.xlsx 3yr_Scenarios (linked by run_all.py)", "[EST] model output", FMT_INT)
    A.add("CA-01", "PNA paying students per year today (all courses) - founder 'under 60'", 30, 50, 60, "students/yr", "Founder 2026-10-16 (canonical CA-01)", "[FOUNDER]", FMT_INT)
    A.add("SOM-BASELINE", "Illustrative PNA baseline revenue today (CA-01 x assumed fees; NOT actual revenue)",
          *[fc[s]["baseline_rev_lakh"] if fc else 0 for s in ("Conservative", "Base", "Aggressive")], "₹ lakh ex-GST",
          "business_forecast.xlsx Unit_Economics baseline (CA-01, CA-02)", "[HYP]", FMT_1DP)

    a = A.r
    # ------------------------------------------------------------------ TAM_India
    ws = wb.create_sheet("TAM_India")
    title(ws, "TAM: India paid makeup-led training tuition (₹ crore / yr, current run-rate)",
          "Re-computes india_market_analysis §8 with live formulas. Definition: learner-paid tuition for makeup-led professional/enthusiast training; excludes govt-funded places, kits, products, salon services.")
    header(ws, 4, ["Step", "Formula", "Low", "Base", "High", "Assumption / note"], widths=[50, 26, 14, 14, 14, 60])
    R = {}

    def line(r, label, formula_fn, fmt, notetxt="", key=None, bold=False):
        ws.cell(row=r, column=1, value=label).font = BOLD if bold else Font()
        for j in range(3):
            c = ws.cell(row=r, column=3 + j, value="=" + formula_fn(j, COLS[j]))
            c.number_format = fmt
            if bold:
                c.font = BOLD
                c.fill = KEY_FILL
        ws.cell(row=r, column=6, value=notetxt).font = NOTE_FONT
        if key:
            R[key] = r

    cs = {0: "low", 1: "base", 2: "high"}
    r = 5
    section(ws, r, "Method 1 - pipeline (council throughput -> paid -> makeup-led)", 6); r += 1
    line(r, "a. Council-certified per year", lambda j, c: a("A-IND-01", cs[j]), FMT_INT, "A-IND-01", "m1a"); ws.cell(row=r, column=2, value="input"); r += 1
    line(r, "b. Total vocational trainees per year", lambda j, c: f"{c}{R['m1a']}/{a('A-IND-02', cs[j])}", FMT_INT, "a / A-IND-02", "m1b"); r += 1
    line(r, "   check: share of KPMG need", lambda j, c: f"{c}{R['m1b']}/{a('X-KPMG-NEED', cs[j])}", FMT_PCT, "b / 679,624", "m1chk"); r += 1
    line(r, "c. Paid share", lambda j, c: a("A-IND-03", cs[j]), FMT_PCT, "A-IND-03", "m1c"); r += 1
    line(r, "d. Paid learners, all beauty", lambda j, c: f"{c}{R['m1b']}*{c}{R['m1c']}", FMT_INT, "b x c", "m1d"); r += 1
    line(r, "e. Makeup-led share", lambda j, c: a("A-IND-04", cs[j]), FMT_PCT, "A-IND-04", "m1e"); r += 1
    line(r, "f. Makeup-led paid learners", lambda j, c: f"{c}{R['m1d']}*{c}{R['m1e']}", FMT_INT, "d x e", "m1f", True); r += 1
    fee = lambda j: (f"{a('A-IND-05a', cs[j])}*{a('A-IND-05b', cs[j])}+{a('A-IND-05c', cs[j])}*{a('A-IND-05d', cs[j])}+"
                     f"{a('A-IND-05e', cs[j])}*{a('A-IND-05f', cs[j])}+{a('A-IND-05g', cs[j])}*{a('A-IND-05h', cs[j])}")
    line(r, "g. Blended fee per makeup-led learner (₹)", lambda j, c: fee(j), FMT_RS, "Σ tier share x tier fee (A-IND-05a..h)", "fee"); r += 1
    line(r, "h. Makeup-led tuition revenue (₹ crore)", lambda j, c: f"{c}{R['m1f']}*{c}{R['fee']}/1E7", FMT_1DP, "f x g / 1e7", "m1h", True); r += 1
    line(r, "i. Non-makeup learners", lambda j, c: f"{c}{R['m1d']}-{c}{R['m1f']}", FMT_INT, "d - f", "m1i"); r += 1
    line(r, "k. Non-makeup revenue (₹ crore)", lambda j, c: f"{c}{R['m1i']}*{a('A-IND-06', cs[j])}/1E7", FMT_1DP, "i x A-IND-06", "m1k"); r += 1
    line(r, "l. All-beauty paid training (₹ crore) - context", lambda j, c: f"{c}{R['m1h']}+{c}{R['m1k']}", FMT_1DP, "h + k", "m1l"); r += 2

    section(ws, r, "Method 2 - supply side (chains x learners x independent multiple)", 6); r += 1
    line(r, "a. Branded chain/franchise centres", lambda j, c: a("A-IND-07", cs[j]), FMT_INT, "A-IND-07", "m2a"); r += 1
    line(r, "b. Learners per centre per year", lambda j, c: a("A-IND-08", cs[j]), FMT_INT, "A-IND-08", "m2b"); r += 1
    line(r, "c. Organised private learners", lambda j, c: f"{c}{R['m2a']}*{c}{R['m2b']}", FMT_INT, "a x b", "m2c"); r += 1
    line(r, "e. Total paid learners, all beauty", lambda j, c: f"{c}{R['m2c']}*(1+{a('A-IND-09', cs[j])})", FMT_INT, "c x (1 + A-IND-09)", "m2e"); r += 1
    line(r, "g. Makeup-led paid learners", lambda j, c: f"{c}{R['m2e']}*{a('A-IND-04', cs[j])}", FMT_INT, "e x A-IND-04 (shared with M1: not independent)", "m2g", True); r += 1
    line(r, "i. Makeup-led tuition revenue (₹ crore)", lambda j, c: f"{c}{R['m2g']}*{c}{R['fee']}/1E7", FMT_1DP, "g x blended fee", "m2i", True); r += 1
    line(r, "j. Non-makeup revenue (₹ crore)", lambda j, c: f"({c}{R['m2e']}-{c}{R['m2g']})*{a('A-IND-06', cs[j])}/1E7", FMT_1DP, "", "m2j"); r += 1
    line(r, "k. All-beauty paid training (₹ crore) - context", lambda j, c: f"{c}{R['m2i']}+{c}{R['m2j']}", FMT_1DP, "", "m2k"); r += 2

    section(ws, r, "Triangulated TAM (mean of methods) - CA-06", 6); r += 1
    line(r, "TAM: makeup-led paid tuition, India (₹ crore/yr)", lambda j, c: f"AVERAGE({c}{R['m1h']},{c}{R['m2i']})", FMT_1DP, "equals CA-06 (₹80 / 390 / 1,370 crore)", "tam", True); r += 1
    line(r, "Makeup-led paid learners, India (mean of methods)", lambda j, c: f"AVERAGE({c}{R['m1f']},{c}{R['m2g']})", FMT_INT, "used for share-of-India arithmetic", "tamL", True); r += 1
    line(r, "Makeup-led learners per 1,000 women aged 17-35", lambda j, c: f"{c}{R['tamL']}/{a('X-W-IND', cs[j])}", FMT_2DP, "national participation intensity", "tamP"); r += 1
    line(r, "Online / short-course layer (₹ crore) - reported separately", lambda j, c: f"{a('A-IND-10a', cs[j])}*{a('A-IND-10b', cs[j])}/1E7", FMT_1DP, "A-IND-10a x 10b", "onl"); r += 1
    line(r, "TAM incl. online layer (₹ crore)", lambda j, c: f"{c}{R['tam']}+{c}{R['onl']}", FMT_1DP, "", "tamOnl", True); r += 1
    line(r, "USD million equivalent (core TAM)", lambda j, c: f"{c}{R['tam']}*1E7/{a('A-IND-16', cs[j])}/1E6", FMT_1DP, "display only", "usd"); r += 1
    line(r, "Context: all-beauty paid training (mean, ₹ crore) - NOT the TAM", lambda j, c: f"AVERAGE({c}{R['m1l']},{c}{R['m2k']})", FMT_1DP, "wider beauty-education context", "allb"); r += 1
    line(r, "Context: core TAM as % of Beauty & Salon sector 2025P", lambda j, c: f"{c}{R['tam']}/{a('X-BS-SECTOR', cs[j])}", FMT_PCT, "shows education ≪ services market", "pctS"); r += 1
    line(r, "Context: 5-state regional slice by KPMG workforce share (₹ crore)", lambda j, c: f"{c}{R['tam']}*{a('A-IND-17', cs[j])}", FMT_1DP, "A-IND-17 proxy; cross-check for SAM_North", "reg765"); r += 1
    note(ws, r + 1, "Methods share A-IND-04 and A-IND-05, so their agreement on order of magnitude is weaker evidence than it looks. Base is a planning number, not a market fact.")
    TAM = R.copy()

    # ------------------------------------------------------------------ Tricity_Recon
    wt = wb.create_sheet("Tricity_Recon")
    title(wt, "Reconciliation: tricity (CA-05) vs India (CA-06)",
          "Why ₹5.7 cr looked like 1.5% of a ₹390 cr India market when the tricity holds ~0.2% of India's young women - and the consistent set we adopt.")
    header(wt, 4, ["Item", "Formula", "Low", "Base", "High", "Note"], widths=[58, 30, 14, 14, 14, 70])
    T = {}

    def tl(r, label, fn, fmt, notetxt="", key=None, bold=False, formula_txt=""):
        wt.cell(row=r, column=1, value=label).font = BOLD if bold else Font()
        wt.cell(row=r, column=2, value=formula_txt).font = NOTE_FONT
        for j in range(3):
            c = wt.cell(row=r, column=3 + j, value="=" + fn(j, COLS[j]))
            c.number_format = fmt
            if bold:
                c.font = BOLD
                c.fill = KEY_FILL
        wt.cell(row=r, column=6, value=notetxt).font = NOTE_FONT
        wt.cell(row=r, column=6).alignment = WRAP
        if key:
            T[key] = r

    TA = "TAM_India!"
    r = 5
    section(wt, r, "0. Demand-intensity indices (computed from [FACT] KPMG state B&W workforce shares)", 6); r += 1
    # rows 6..12 - keep the KPMG index at row 12 because A-ECO-01 base links to Tricity_Recon!D12
    tl(r, "Chandigarh B&W intensity (state share / India share)", lambda j, c: f"{a('X-BW-CHD', cs[j])}/{a('X-BW-IND', cs[j])}", FMT_2DP, "6.16% / 2.41%", "iChd"); r += 1
    tl(r, "Punjab intensity (applied to Mohali)", lambda j, c: f"{a('X-BW-PB', cs[j])}/{a('X-BW-IND', cs[j])}", FMT_2DP, "", "iPb"); r += 1
    tl(r, "Haryana intensity (applied to Panchkula)", lambda j, c: f"{a('X-BW-HR', cs[j])}/{a('X-BW-IND', cs[j])}", FMT_2DP, "", "iHr"); r += 1
    tl(r, "Women weights: Chandigarh / Mohali / Panchkula (000)", lambda j, c: f"{a('X-W-CHD', cs[j])}+{a('X-W-MOH', cs[j])}+{a('X-W-PKL', cs[j])}", FMT_INT, "sum = 471k", "wsum"); r += 1
    tl(r, "Affluence weight (A-CHD-06) - cross-check", lambda j, c: a("A-CHD-06", cs[j]), FMT_2DP, "income-based; agrees with KPMG index within ~3% at base", "aff"); r += 1
    tl(r, "(spacer)", lambda j, c: "0", FMT_GEN, ""); r += 1
    assert r == 12
    tl(r, "KPMG B&W intensity index, tricity (women-weighted)",
       lambda j, c: f"({a('X-W-CHD', cs[j])}*{c}{T['iChd']}+{a('X-W-MOH', cs[j])}*{c}{T['iPb']}+{a('X-W-PKL', cs[j])}*{c}{T['iHr']})/{c}{T['wsum']}",
       FMT_2DP, "Feeds A-ECO-01 base (1.90)", "kpmgT", True); r += 2

    section(wt, r, "1. As published (scopes differ)", 6); r += 1
    tl(r, "CA-05 tricity learners = women x participation / (1 - outstation)",
       lambda j, c: f"{a('A-CHD-02', cs[j])}*1000*{a('A-CHD-03', cs[j])}/(1-{a('A-CHD-04', cs[j])})", FMT_INT,
       "chandigarh_market_analysis §8.1: 593 / 1,256 / 2,229. Scope: makeup + beauty-artistry; participation calibrated to bridal-MUA pipeline", "ca05L")
    r += 1
    tl(r, "CA-05 tricity revenue (₹ crore)", lambda j, c: f"{c}{T['ca05L']}*{a('A-CHD-05', cs[j])}/1E7", FMT_2DP, "1.8 / 5.7 / 13.4", "ca05R", True); r += 1
    tl(r, "CA-06 India TAM (₹ crore)", lambda j, c: f"{TA}{c}{TAM['tam']}", FMT_1DP, "80 / 390 / 1,370", "ca06"); r += 1
    tl(r, "Apparent tricity share of India (CA-05 / CA-06)", lambda j, c: f"{c}{T['ca05R']}/{c}{T['ca06']}", FMT_PCT2, "the 'open reconciliation': ~1.5% at base", "appShare", True); r += 1
    tl(r, "Tricity share of India women 17-35", lambda j, c: f"{a('A-CHD-02', cs[j])}/{a('X-W-IND', cs[j])}", FMT_PCT2, "~0.21%", "popShare"); r += 2

    section(wt, r, "2. Method A - top-down from India TAM (national participation x tricity intensity)", 6); r += 1
    tl(r, "India makeup-led paid learners", lambda j, c: f"{TA}{c}{TAM['tamL']}", FMT_INT, "TAM_India mean of methods", "indL"); r += 1
    tl(r, "x tricity population share", lambda j, c: f"{c}{T['popShare']}", FMT_PCT2, "", "pA"); r += 1
    tl(r, "x demand-intensity index (A-ECO-01)", lambda j, c: a("A-ECO-01", cs[j]), FMT_2DP, "B&W workforce / affluence", "intA"); r += 1
    tl(r, "x urban-concentration factor = (uT*k + 1-uT)/(uI*k + 1-uI)",
       lambda j, c: f"({a('A-ECO-03', cs[j])}*{a('A-ECO-02', cs[j])}+1-{a('A-ECO-03', cs[j])})/({a('A-ECO-04', cs[j])}*{a('A-ECO-02', cs[j])}+1-{a('A-ECO-04', cs[j])})",
       FMT_2DP, "A-ECO-02..04. May overlap partly with the intensity index (both rise with urbanisation) - flagged", "urbA"); r += 1
    tl(r, "/ (1 - outstation inflow)", lambda j, c: f"1/(1-{a('A-CHD-04', cs[j])})", FMT_2DP, "tricity counts students who come from home states", "inA"); r += 1
    tl(r, "Tricity share of India makeup-led learners (product)", lambda j, c: f"{c}{T['pA']}*{c}{T['intA']}*{c}{T['urbA']}*{c}{T['inA']}", FMT_PCT2, "", "shA", True); r += 1
    tl(r, "Method A learners", lambda j, c: f"{c}{T['indL']}*{c}{T['shA']}", FMT_INT, "", "mA", True); r += 1
    tl(r, "Method A revenue at tricity fee (₹ crore)", lambda j, c: f"{c}{T['mA']}*{a('A-CHD-05', cs[j])}/1E7", FMT_2DP, "tricity fee (A-CHD-05) < India blended fee (no T4 premium tier locally)", "mAR"); r += 2

    section(wt, r, "3. Method B - CA-05 bottom-up restated to the TAM scope (makeup-led only)", 6); r += 1
    tl(r, "CA-05 learners (makeup + beauty-artistry)", lambda j, c: f"{c}{T['ca05L']}", FMT_INT, "", "bL0"); r += 1
    tl(r, "x makeup-led share (A-ECO-06)", lambda j, c: a("A-ECO-06", cs[j]), FMT_PCT, "scope harmonisation", "bSh"); r += 1
    tl(r, "Method B learners (makeup-led)", lambda j, c: f"{c}{T['bL0']}*{c}{T['bSh']}", FMT_INT, "", "mB", True); r += 1
    tl(r, "Method B revenue (₹ crore)", lambda j, c: f"{c}{T['mB']}*{a('A-CHD-05', cs[j])}/1E7", FMT_2DP, "", "mBR"); r += 2

    section(wt, r, "4. Method C - supply-side count of tricity providers (cross-check)", 6); r += 1
    tl(r, "Chain/brand centres x makeup learners", lambda j, c: f"{a('A-ECO-07', cs[j])}*{a('A-ECO-08', cs[j])}", FMT_INT, "A-ECO-07 x 08", "c1"); r += 1
    tl(r, "Independent institutes x learners", lambda j, c: f"{a('A-ECO-09', cs[j])}*{a('A-ECO-10', cs[j])}", FMT_INT, "A-ECO-09 x 10", "c2"); r += 1
    tl(r, "Founder-led academies x learners", lambda j, c: f"{a('A-ECO-11', cs[j])}*{a('A-ECO-12', cs[j])}", FMT_INT, "A-ECO-11 x 12", "c3"); r += 1
    tl(r, "Tutors (UrbanPro listings x active learners per listing)", lambda j, c: f"{a('A-ECO-13', cs[j])}*{a('A-ECO-14', cs[j])}", FMT_INT, "A-ECO-13 x 14", "c4"); r += 1
    tl(r, "Method C learners", lambda j, c: f"SUM({c}{T['c1']}:{c}{T['c4']})", FMT_INT, "", "mC", True); r += 1
    tl(r, "Method C revenue (₹ crore)", lambda j, c: f"{c}{T['mC']}*{a('A-CHD-05', cs[j])}/1E7", FMT_2DP, "", "mCR"); r += 2

    section(wt, r, "5. ADOPTED CONSISTENT SET - tricity makeup-led paid training (planning band = spread of the three methods' base cases)", 6); r += 1
    tl(r, "Reconciled tricity learners (Low = min, Base = mean, High = max of method base cases)",
       lambda j, c: ["MIN", "AVERAGE", "MAX"][j] + f"($D${T['mA']},$D${T['mB']},$D${T['mC']})", FMT_INT,
       "Method disagreement is the dominant uncertainty; each method's own low/high compounds extreme inputs (shown below)", "recL", True); r += 1
    tl(r, "Reconciled tricity revenue (₹ crore, quoted fees)", lambda j, c: f"{c}{T['recL']}*{a('A-CHD-05', cs[j])}/1E7", FMT_2DP,
       "learners x A-CHD-05 (₹30k / 45k / 60k)", "recR", True); r += 1
    tl(r, "  of which outstation (inflow) learners", lambda j, c: f"{c}{T['recL']}*{a('A-CHD-04', cs[j])}", FMT_INT, "", "recOut"); r += 1
    tl(r, "Reconciled tricity revenue as share of India BASE TAM", lambda j, c: f"{c}{T['recR']}/$D${T['ca06']}", FMT_PCT2, "Base ≈0.84% vs 1.46% apparent (CA-05/CA-06)", "recSh", True); r += 1
    tl(r, "Outer bounds: min of method lows / max of method highs (learners)",
       lambda j, c: ["MIN($C${0},$C${1},$C${2})".format(T['mA'], T['mB'], T['mC']), f"$D${T['recL']}",
                     "MAX($E${0},$E${1},$E${2})".format(T['mA'], T['mB'], T['mC'])][j], FMT_INT, "for transparency only", "outer"); r += 2

    section(wt, r, "6. Decomposing the base-case gap (participation per 1,000 women 17-35)", 6); r += 1
    tl(r, "CA-05 tricity learners per 1,000 women", lambda j, c: f"{c}{T['ca05L']}/{a('A-CHD-02', cs[j])}", FMT_2DP, "", "gT"); r += 1
    tl(r, "India makeup-led learners per 1,000 women", lambda j, c: f"{TA}{c}{TAM['tamP']}", FMT_2DP, "", "gI"); r += 1
    tl(r, "Gap ratio to explain", lambda j, c: f"{c}{T['gT']}/{c}{T['gI']}", FMT_X, "", "gap", True); r += 1
    tl(r, "  explained by scope (CA-05 includes non-makeup artistry) = 1/A-ECO-06", lambda j, c: f"1/{c}{T['bSh']}", FMT_X, "", "g1"); r += 1
    tl(r, "  explained by outstation inflow = 1/(1-A-CHD-04)", lambda j, c: f"{c}{T['inA']}", FMT_X, "", "g2"); r += 1
    tl(r, "  explained by demand intensity (A-ECO-01)", lambda j, c: f"{c}{T['intA']}", FMT_X, "", "g3"); r += 1
    tl(r, "  explained by urban concentration", lambda j, c: f"{c}{T['urbA']}", FMT_X, "", "g4"); r += 1
    tl(r, "  product of explained factors", lambda j, c: f"{c}{T['g1']}*{c}{T['g2']}*{c}{T['g3']}*{c}{T['g4']}", FMT_X, "", "gEx", True); r += 1
    tl(r, "  unexplained residual", lambda j, c: f"{c}{T['gap']}/{c}{T['gEx']}", FMT_X,
       "Residual = local evidence (CA-05 pipeline, Method C supply count) sits above the national top-down. Read: CA-06 base is conservative for urban North India and/or tricity is a makeup-academy hot-spot.", "gRes", True); r += 1
    tl(r, "Residual retained in adopted base (reconciled / Method A)", lambda j, c: f"$D${T['recL']}/$D${T['mA']}", FMT_X, "the planning set keeps ~half of the residual (geometric middle)", "gKeep"); r += 2

    section(wt, r, "7. Resolution (documented in assumptions_register.md and assumptions_and_methodology.md)", 6); r += 1
    res = [
        "1. Scope: one definition everywhere - learner-paid, private, makeup-led training tuition at quoted fees (GST-inclusive where quoted); excludes non-makeup beauty courses, govt-funded places, kits, products, services.",
        "2. India TAM (CA-06) is retained unchanged: ₹80 / 390 / 1,370 crore (+ online layer ₹3 / 23 / 120 crore shown separately).",
        "3. CA-05 (₹5.7 crore) is NOT a makeup-led number: it counts makeup + beauty-artistry learners at a diploma fee. Restated to makeup-led scope it is Method B.",
        "4. Adopted tricity makeup-led market = mean of three methods (top-down, CA-05 restated, supply count): ≈₹3.3 crore and ≈730 learners at base (band ₹1.4-5.3 crore).",
        "5. This is ≈0.8-0.9% of the India base TAM - above the 0.21% population share because of beauty-workforce intensity (1.9x), urban concentration (1.6x) and outstation inflow (1.33x); the remaining ~1.5x is local evidence the national model does not see.",
        "6. Consequence: tricity alone caps a single centre at roughly ₹0.5-0.8 crore of tuition at a 15-25% share - growth must come from regional, hybrid, CPD and online layers (SAM_North).",
    ]
    for t in res:
        wt.cell(row=r, column=1, value=t).alignment = WRAP
        wt.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        wt.row_dimensions[r].height = 30
        r += 1

    # ------------------------------------------------------------------ SAM_North
    wsm = wb.create_sheet("SAM_North")
    title(wsm, "SAM: North India market serviceable by a Chandigarh-led offline/hybrid academy (₹ crore / yr, quoted fees)",
          "Five components built on the same participation chain as the reconciled tricity number; outstation learners already inside the tricity figure are netted out of the regional component.")
    header(wsm, 4, ["Item", "Formula", "Low", "Base", "High", "Note"], widths=[60, 30, 14, 14, 14, 64])
    S = {}

    def sl(r, label, fn, fmt, notetxt="", key=None, bold=False):
        wsm.cell(row=r, column=1, value=label).font = BOLD if bold else Font()
        for j in range(3):
            c = wsm.cell(row=r, column=3 + j, value="=" + fn(j, COLS[j]))
            c.number_format = fmt
            if bold:
                c.font = BOLD
                c.fill = KEY_FILL
        wsm.cell(row=r, column=6, value=notetxt).font = NOTE_FONT
        if key:
            S[key] = r

    TR = "Tricity_Recon!"
    r = 5
    section(wsm, r, "Rest-of-region intensity (PB/HR/HP/J&K excl. tricity)", 6); r += 1
    sl(r, "Σ women x state B&W intensity",
       lambda j, c: (f"({a('X-W-PB', cs[j])}*{a('X-BW-PB', cs[j])}+{a('X-W-HR', cs[j])}*{a('X-BW-HR', cs[j])}+"
                     f"{a('X-W-HP', cs[j])}*{a('X-BW-HP', cs[j])}+{a('X-W-JK', cs[j])}*{a('X-BW-JK', cs[j])})/{a('X-BW-IND', cs[j])}"), FMT_INT, "", "num"); r += 1
    sl(r, "Σ women (000)", lambda j, c: f"{a('X-W-PB', cs[j])}+{a('X-W-HR', cs[j])}+{a('X-W-HP', cs[j])}+{a('X-W-JK', cs[j])}", FMT_INT, "equals X-W-ROR 12,376", "den"); r += 1
    sl(r, "(spacer)", lambda j, c: "0", FMT_GEN, ""); r += 1
    assert r == 9
    sl(r, "KPMG intensity index, rest of region", lambda j, c: f"{c}{S['num']}/{c}{S['den']}", FMT_2DP, "Feeds A-ECO-05 base (1.30)", "iRoR", True); r += 2

    section(wsm, r, "SAM-1 Tricity (reconciled, Tricity_Recon §5)", 6); r += 1
    sl(r, "Tricity makeup-led learners", lambda j, c: f"{TR}{c}{T['recL']}", FMT_INT, "", "s1L"); r += 1
    sl(r, "SAM-1 revenue (₹ crore)", lambda j, c: f"{TR}{c}{T['recR']}", FMT_2DP, "", "s1", True); r += 2

    section(wsm, r, "SAM-2 Regional destination learners (offline, travel to Chandigarh)", 6); r += 1
    sl(r, "Rest-of-region makeup-led learners", lambda j, c: f"{TA}{c}{TAM['tamL']}*{a('X-W-ROR', cs[j])}/{a('X-W-IND', cs[j])}*{a('A-ECO-05', cs[j])}", FMT_INT,
       "India learners x population share x A-ECO-05", "rorL"); r += 1
    sl(r, "x reachable share (A-ECO-15)", lambda j, c: a("A-ECO-15", cs[j]), FMT_PCT, "", "reach"); r += 1
    sl(r, "Reachable regional learners", lambda j, c: f"{c}{S['rorL']}*{c}{S['reach']}", FMT_INT, "", "reachL"); r += 1
    sl(r, "less: already counted as tricity outstation inflow", lambda j, c: f"{TR}{c}{T['recOut']}", FMT_INT, "avoid double count", "less"); r += 1
    sl(r, "Net incremental regional learners", lambda j, c: f"MAX(0,{c}{S['reachL']}-{c}{S['less']})", FMT_INT, "", "s2L"); r += 1
    sl(r, "SAM-2 revenue (₹ crore)", lambda j, c: f"{c}{S['s2L']}*{a('A-CHD-05', cs[j])}/1E7", FMT_2DP, "at tricity fee", "s2", True); r += 2

    section(wsm, r, "SAM-3 Hybrid (online theory + short Chandigarh practical) for learners beyond offline reach", 6); r += 1
    sl(r, "Non-reachable regional learners", lambda j, c: f"{c}{S['rorL']}-{c}{S['reachL']}", FMT_INT, "", "nrL"); r += 1
    sl(r, "x hybrid-willing share (A-ECO-16)", lambda j, c: a("A-ECO-16", cs[j]), FMT_PCT, "", "hyb"); r += 1
    sl(r, "Hybrid learners", lambda j, c: f"{c}{S['nrL']}*{c}{S['hyb']}", FMT_INT, "", "s3L"); r += 1
    sl(r, "SAM-3 revenue (₹ crore)", lambda j, c: f"{c}{S['s3L']}*{a('A-ECO-17', cs[j])}/1E7", FMT_2DP, "A-ECO-17 fee", "s3", True); r += 2

    section(wsm, r, "SAM-4 Working-artist CPD / masterclasses (tricity + reachable region) - scope extension, see note", 6); r += 1
    sl(r, "Working makeup artists in reach", lambda j, c: f"{a('A-ECO-18', cs[j])}+{a('A-ECO-19', cs[j])}", FMT_INT, "A-ECO-18 + 19", "s4A"); r += 1
    sl(r, "Paying CPD participants per year", lambda j, c: f"{c}{S['s4A']}*{a('A-ECO-20', cs[j])}", FMT_INT, "", "s4L"); r += 1
    sl(r, "SAM-4 revenue (₹ crore)", lambda j, c: f"{c}{S['s4L']}*{a('A-ECO-21', cs[j])}/1E7", FMT_2DP,
       "CPD by working pros is largely outside the TAM's entrant-based learner count; shown separately so it does not inflate TAM comparisons", "s4", True); r += 2

    section(wsm, r, "SAM-5 Online short-course layer reachable with Hindi/Punjabi founder content", 6); r += 1
    sl(r, "SAM-5 revenue (₹ crore)", lambda j, c: f"{TA}{c}{TAM['onl']}*{a('A-ECO-22', cs[j])}", FMT_2DP, "India online layer x A-ECO-22", "s5", True); r += 2

    section(wsm, r, "SAM totals", 6); r += 1
    sl(r, "SAM core: offline + hybrid + CPD (SAM-1..4), ₹ crore", lambda j, c: f"{c}{S['s1']}+{c}{S['s2']}+{c}{S['s3']}+{c}{S['s4']}", FMT_2DP, "headline SAM", "core", True); r += 1
    sl(r, "SAM online layer (SAM-5), ₹ crore", lambda j, c: f"{c}{S['s5']}", FMT_2DP, "low-ticket, crowded (Orane, Airblack, YouTube)", "onl"); r += 1
    sl(r, "SAM total, ₹ crore", lambda j, c: f"{c}{S['core']}+{c}{S['onl']}", FMT_2DP, "", "tot", True); r += 1
    sl(r, "Training learners in SAM-1..3 (excl. CPD and online)", lambda j, c: f"{c}{S['s1L']}+{c}{S['s2L']}+{c}{S['s3L']}", FMT_INT, "", "learn"); r += 1
    sl(r, "SAM total / TAM incl. online", lambda j, c: f"{c}{S['tot']}/{TA}{c}{TAM['tamOnl']}", FMT_PCT, "", "samTam"); r += 1
    sl(r, "Cross-check: 5-state makeup-led market bottom-up (tricity + all rest-of-region learners x tricity fee)",
       lambda j, c: f"{c}{S['s1']}+{c}{S['rorL']}*{a('A-CHD-05', cs[j])}/1E7", FMT_2DP, "compare with next row", "chk1"); r += 1
    sl(r, "Cross-check: TAM x 7.65% workforce proxy, restated at tricity/India fee ratio",
       lambda j, c: f"{TA}{c}{TAM['reg765']}*{a('A-CHD-05', cs[j])}/{TA}{c}{TAM['fee']}", FMT_2DP, "two routes should be the same order of magnitude", "chk2"); r += 1

    # ------------------------------------------------------------------ SOM_PNA
    wso = wb.create_sheet("SOM_PNA")
    title(wso, "SOM: PNA's realistic Year-3 revenue and implied share of each SAM component",
          "SOM = Year-3 revenue from business_forecast.xlsx 3yr_Scenarios (Low=Conservative, Base, High=Aggressive), converted to quoted (GST-inclusive) prices to match market sizing.")
    header(wso, 4, ["Item", "Formula", "Low (Cons.)", "Base", "High (Aggr.)", "Note"], widths=[60, 30, 14, 14, 14, 64])
    O = {}

    def ol(r, label, fn, fmt, notetxt="", key=None, bold=False):
        wso.cell(row=r, column=1, value=label).font = BOLD if bold else Font()
        for j in range(3):
            c = wso.cell(row=r, column=3 + j, value="=" + fn(j, COLS[j]))
            c.number_format = fmt
            if bold:
                c.font = BOLD
                c.fill = KEY_FILL
        wso.cell(row=r, column=6, value=notetxt).font = NOTE_FONT
        if key:
            O[key] = r

    SN = "SAM_North!"
    g = lambda j: f"(1+{a('A-ECO-23', cs[j])})"
    r = 5
    pairs = [("SOM-TRI", "s1", "Tricity flagship (SAM-1)"), ("SOM-REG", "s2", "Regional destination (SAM-2)"),
             ("SOM-HYB", "s3", "Hybrid (SAM-3)"), ("SOM-CPD", "s4", "Working-artist CPD (SAM-4)"), ("SOM-ONL", "s5", "Online (SAM-5)")]
    section(wso, r, "SOM by component (₹ crore, quoted incl. GST) and implied share", 6); r += 1
    somrows = []
    for sid, skey, lab in pairs:
        ol(r, f"{lab}: PNA Y3 revenue", lambda j, c, sid=sid: f"{a(sid, cs[j])}*{g(j)}/100", FMT_2DP, f"{sid} x (1+GST) / 100", "v" + skey); somrows.append(r); r += 1
        ol(r, f"{lab}: implied share of SAM component", lambda j, c, skey=skey: f"IFERROR({c}{O['v' + skey]}/{SN}{c}{S[skey]},0)", FMT_PCT, "", "sh" + skey, True); r += 1
    section(wso, r, "Learner-share test (primary: PNA charges ~1.8x the tricity average fee, so revenue share overstates position)", 6); r += 1
    ol(r, "PNA Y3 tricity career learners / tricity makeup-led learners (SAM-1)", lambda j, c: f"{a('SOM-LRN-TRI', cs[j])}/{SN}{c}{S['s1L']}", FMT_PCT, "Chandigarh agent: 10-30% is the plausible single-centre band", "lsh1", True); r += 1
    ol(r, "PNA Y3 out-of-town career learners / net regional reachable learners (SAM-2)", lambda j, c: f"IFERROR({a('SOM-LRN-REG', cs[j])}/({SN}{c}{S['s2L']}+{TR}{c}{T['recOut']}),0)", FMT_PCT, "denominator incl. inflow already in tricity", "lsh2", True); r += 1
    ol(r, "PNA Y3 hybrid learners / hybrid-willing learners (SAM-3)", lambda j, c: f"IFERROR({a('SOM-LRN-HYB', cs[j])}/{SN}{c}{S['s3L']},0)", FMT_PCT, "", "lsh3", True); r += 1
    ol(r, "PNA Y3 tricity career learners / CA-05 tricity learners (BM test: flag if >10% without regional engine)", lambda j, c: f"{a('SOM-LRN-TRI', cs[j])}/{TR}{c}{T['ca05L']}", FMT_PCT, "", "lsh4"); r += 1
    ol(r, "SOM training total (₹ crore, quoted)", lambda j, c: "+".join(f"{c}{x}" for x in somrows), FMT_2DP, "", "som", True); r += 1
    ol(r, "SOM / SAM core (excl. online)", lambda j, c: f"({c}{O['som']}-{c}{O['vs5']})/{SN}{c}{S['core']}", FMT_PCT, "", "somCore", True); r += 1
    ol(r, "SOM / SAM total", lambda j, c: f"{c}{O['som']}/{SN}{c}{S['tot']}", FMT_PCT, "", "somSam"); r += 1
    ol(r, "SOM / TAM incl. online", lambda j, c: f"{c}{O['som']}/{TA}{c}{TAM['tamOnl']}", FMT_PCT2, "benchmark: Mao Geping ≈1% of China makeup training (S-INT-01)", "somTam"); r += 1
    ol(r, "Adjacent revenue outside training market (₹ crore, quoted)", lambda j, c: f"{a('SOM-ADJ', cs[j])}*{g(j)}/100", FMT_2DP, "kits, services desk, partnerships - not counted in SOM share", "adj"); r += 2
    section(wso, r, "Today (illustrative baseline - CA-01 x assumed fees, NOT actual revenue)", 6); r += 1
    ol(r, "PNA illustrative baseline (₹ crore, quoted)", lambda j, c: f"{a('SOM-BASELINE', cs[j])}*{g(j)}/100", FMT_2DP, "CA-01 50 students x CA-02 fee mix", "base"); r += 1
    ol(r, "Implied baseline share of tricity SAM-1", lambda j, c: f"{c}{O['base']}/{SN}{c}{S['s1']}", FMT_PCT, "sanity: ~5-10% for an established founder-led studio-academy", "baseSh", True); r += 1
    ol(r, "Y3 SOM multiple of today", lambda j, c: f"IFERROR({c}{O['som']}/{c}{O['base']},0)", FMT_X, "", "mult"); r += 1
    ol(r, "ILLUSTRATIVE: today's PNA share of tricity makeup-led learners (CA-01 / reconciled learners; Low = CA-01 low / market high, High = CA-01 high / market low)",
       lambda j, c: [f"{a('CA-01', 'low')}/{TR}$E${T['recL']}", f"{a('CA-01', 'base')}/{TR}$D${T['recL']}", f"{a('CA-01', 'high')}/{TR}$C${T['recL']}"][j], FMT_PCT,
       "All PNA students (career + short) over all makeup-led learners (incl. short-course). Not quoted as a fact until CA-01 is confirmed (V1)", "todaySh", True); r += 1
    note(wso, r + 1, "Share tests: a single centre at >25% of tricity, >30% of CPD or >5% of online would be a red flag (chandigarh_market_analysis §8.2; international_benchmarks §2).")

    # ------------------------------------------------------------------ Sizing_Tree
    wtr = wb.create_sheet("Sizing_Tree")
    title(wtr, "Market-sizing tree (₹ crore / yr unless stated)", "Never present the beauty products/services market as the education market (see Context_NotTAM).")
    header(wtr, 4, ["Node", "", "Low", "Base", "High", "Selected case", "Definition / source"], widths=[56, 4, 13, 13, 13, 14, 66])
    tree = [
        ("TAM  India: paid makeup-led training tuition", TA, TAM['tam'], FMT_1DP, "CA-06; TAM_India (mean of pipeline and supply methods)"),
        ("  + India online / short-course layer", TA, TAM['onl'], FMT_1DP, "A-IND-10; reported separately"),
        ("  (context) 5-state North slice by KPMG workforce share", TA, TAM['reg765'], FMT_1DP, "A-IND-17 x TAM; not used directly"),
        ("SAM  North India serviceable by a Chandigarh-led academy (total)", SN, S['tot'], FMT_2DP, "SAM_North"),
        ("  SAM core (offline + hybrid + CPD)", SN, S['core'], FMT_2DP, "headline SAM"),
        ("    SAM-1 Tricity makeup-led (reconciled)", SN, S['s1'], FMT_2DP, "Tricity_Recon §5 (replaces CA-05 for planning)"),
        ("    SAM-2 Regional destination (net of inflow)", SN, S['s2'], FMT_2DP, "PB/HR/HP/J&K reachable offline"),
        ("    SAM-3 Hybrid beyond offline reach", SN, S['s3'], FMT_2DP, ""),
        ("    SAM-4 Working-artist CPD / masterclasses", SN, S['s4'], FMT_2DP, "scope extension (repeat buyers)"),
        ("  SAM-5 Online layer (North-reachable)", SN, S['s5'], FMT_2DP, ""),
        ("SOM  PNA Year-3 revenue (quoted, training only)", "SOM_PNA!", O['som'], FMT_2DP, "business_forecast 3yr_Scenarios (Cons/Base/Aggr)"),
        ("  SOM as % of SAM core (excl. online)", "SOM_PNA!", O['somCore'], FMT_PCT, ""),
        ("  PNA learner share of tricity makeup-led learners (Y3)", "SOM_PNA!", O['lsh1'], FMT_PCT, "primary share test"),
        ("  SOM as % of TAM incl. online", "SOM_PNA!", O['somTam'], FMT_PCT2, ""),
        ("Today  PNA illustrative baseline (NOT actual revenue)", "SOM_PNA!", O['base'], FMT_2DP, "CA-01 x CA-02 assumptions"),
        ("  Illustrative: today's PNA share of tricity makeup-led learners", "SOM_PNA!", O['todaySh'], FMT_PCT, "CA-01 / reconciled learners; to confirm with V1"),
        ("Learners: India makeup-led paid learners", TA, TAM['tamL'], FMT_INT, ""),
        ("Learners: tricity makeup-led (reconciled)", TR, T['recL'], FMT_INT, ""),
    ]
    rr = 5
    for lab, sh, row, fmt, d in tree:
        c = wtr.cell(row=rr, column=1, value=lab)
        if not lab.startswith(" "):
            c.font = BOLD
        for j in range(3):
            cc = wtr.cell(row=rr, column=3 + j, value=f"={sh}{COLS[j]}{row}")
            cc.number_format = fmt
        cc = wtr.cell(row=rr, column=6, value=f"=INDEX(C{rr}:E{rr},Assumptions!$C$4)")
        cc.number_format = fmt
        cc.fill = KEY_FILL
        wtr.cell(row=rr, column=7, value=d).font = NOTE_FONT
        rr += 1
    note(wtr, rr + 1, "Selected case follows the selector on Assumptions!C3 (Low / Base / High).")

    # ------------------------------------------------------------------ Context_NotTAM
    wc = wb.create_sheet("Context_NotTAM")
    title(wc, "Context only - these are NOT the education market and must never be presented as TAM")
    header(wc, 4, ["Market", "Figure", "Unit", "Source", "Why excluded"], widths=[44, 18, 16, 30, 60])
    ctx = [("Beauty & Wellness sector (products + services) 2025P", 276999, "₹ crore", "S-IND-01 Table 1", "Consumer spend on products and services"),
           ("Beauty & Salon sub-sector 2025P", 76141, "₹ crore", "S-IND-01 Table 1", "Salon services spend"),
           ("Makeup products (2024)", 1.52, "USD bn", "S-IND-51", "Product sales"),
           ("Beauty salon services (2025)", 12.6, "USD bn", "S-IND-22", "Services; vendor model"),
           ("Bridal makeup service spend (cross-check, base)", 2880, "₹ crore", "india_market_analysis §8.4", "Services demand for MUAs, not tuition"),
           ("Weddings: peak-season business 2025", 650000, "₹ crore", "S-IND-10 (CAIT)", "Wedding industry"),
           ("All-beauty paid training (mean, base) - wider education context", f"=TAM_India!D{TAM['allb']}", "₹ crore", "TAM_India", "Includes non-makeup beauty courses; TAM is the makeup-led subset")]
    for i, row in enumerate(ctx):
        for j, v in enumerate(row):
            c = wc.cell(row=5 + i, column=1 + j, value=v)
            if j == 1:
                c.number_format = FMT_1DP
    wc.cell(row=13, column=1, value="TAM (makeup-led paid tuition, base) for comparison").font = BOLD
    wc.cell(row=13, column=2, value=f"=TAM_India!D{TAM['tam']}").number_format = FMT_1DP

    # ------------------------------------------------------------------ Checks
    wk = wb.create_sheet("Checks")
    title(wk, "Integrity checks (TRUE = pass)")
    header(wk, 4, ["Check", "Result", "Detail"], widths=[70, 12, 50])
    checks = [
        ("TAM_India Method 1 base reproduces research (₹522.1 crore ±1)", f"=ABS(TAM_India!D{TAM['m1h']}-522.1)<1", ""),
        ("TAM base reproduces CA-06 (₹387.7 crore ±1)", f"=ABS(TAM_India!D{TAM['tam']}-387.7)<1", ""),
        ("CA-05 base reproduces research (₹5.65 crore ±0.1)", f"=ABS(Tricity_Recon!D{T['ca05R']}-5.65)<0.1", ""),
        ("Reconciled tricity base lies between Method A and Method B", f"=AND(Tricity_Recon!D{T['recL']}>=Tricity_Recon!D{T['mA']},Tricity_Recon!D{T['recL']}<=Tricity_Recon!D{T['mB']})", ""),
        ("Low <= Base <= High for reconciled tricity revenue", f"=AND(Tricity_Recon!C{T['recR']}<=Tricity_Recon!D{T['recR']},Tricity_Recon!D{T['recR']}<=Tricity_Recon!E{T['recR']})", ""),
        ("SAM total < TAM incl. online (all cases)", f"=AND(SAM_North!C{S['tot']}<TAM_India!C{TAM['tamOnl']},SAM_North!D{S['tot']}<TAM_India!D{TAM['tamOnl']},SAM_North!E{S['tot']}<TAM_India!E{TAM['tamOnl']})", ""),
        ("SOM < SAM total (all cases)", f"=AND(SOM_PNA!C{O['som']}<SAM_North!C{S['tot']},SOM_PNA!D{O['som']}<SAM_North!D{S['tot']},SOM_PNA!E{O['som']}<SAM_North!E{S['tot']})", ""),
        ("Base SOM learner share of tricity <= 25% (revenue share shown separately)", f"=SOM_PNA!D{O['lsh1']}<=0.25", ""),
        ("KPMG tricity intensity within 10% of affluence weight (cross-check)", f"=ABS(Tricity_Recon!D{T['kpmgT']}/Tricity_Recon!D{T['aff']}-1)<0.1", ""),
    ]
    for i, (lab, f, d) in enumerate(checks):
        wk.cell(row=5 + i, column=1, value=lab)
        wk.cell(row=5 + i, column=2, value=f).font = BOLD
        wk.cell(row=5 + i, column=3, value=d)
    wk.cell(row=5 + len(checks) + 1, column=1, value="All checks pass").font = BOLD
    wk.cell(row=5 + len(checks) + 1, column=2, value=f"=AND(B5:B{4 + len(checks)})").font = BOLD

    # ------------------------------------------------------------------ README
    title(readme, "PNA market sizing model - TAM / SAM / SOM with tricity-India reconciliation",
          "Prepared by the Economics & Market Sizing agent, 2026-10-16. Build script: economics/build_scripts/build_market_sizing.py")
    lines = [
        ("Purpose", "Size India's paid makeup-led training market (TAM), the North India slice a Chandigarh-led offline/hybrid academy can serve (SAM) and PNA's realistic Year-3 revenue (SOM); reconcile CA-05 (tricity) with CA-06 (India)."),
        ("Unit", "₹ crore per year, nominal, current run-rate (≈FY2026-27). Market figures at quoted fees (GST-inclusive where quoted). PNA forecast revenue is ex-GST and is grossed up by 18% (A-ECO-23) for SOM shares."),
        ("Sheets", "Assumptions (all inputs, IDs, sources, tags, Low/Base/High + selector) -> TAM_India -> Tricity_Recon -> SAM_North -> SOM_PNA -> Sizing_Tree (summary) ; Context_NotTAM ; Checks"),
        ("Colour code", "Blue = input; black = formula; yellow = key output. No hard-coded results in calculation sheets."),
        ("Selector", "Assumptions!C3 picks Low/Base/High for the 'Selected' column and the Sizing_Tree 'Selected case' column. Calculation sheets show all three cases side by side."),
        ("Reconciliation", "See Tricity_Recon §5-7: CA-06 retained; CA-05 restated to makeup-led scope; adopted tricity = mean of three methods ≈ ₹3.3 crore base."),
        ("Guardrail", "The beauty products / salon services / wedding markets are listed only in Context_NotTAM and are never the education TAM."),
        ("Limitations", "No primary enrolment data; makeup-led share (A-IND-04), fee mix (A-IND-05), participation (A-CHD-03) and reach shares (A-ECO-15/16) are hypotheses. Validate via mystery shopping, competitor intake checks and PNA enquiry-origin data."),
    ]
    for i, (k, v) in enumerate(lines):
        readme.cell(row=4 + i, column=1, value=k).font = BOLD
        c = readme.cell(row=4 + i, column=2, value=v)
        c.alignment = WRAP
        readme.row_dimensions[4 + i].height = 45
    setw(readme, {"A": 18, "B": 120})

    wb.save(OUT)
    return OUT, dict(TAM=TAM, T=T, S=S, O=O)


if __name__ == "__main__":
    out, refs = build()
    json.dump(refs, open(os.path.join(os.path.dirname(__file__), "_sizing_refs.json"), "w"), indent=1)
    print(out)
