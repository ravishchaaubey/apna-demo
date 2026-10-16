"""Build economics/funnel_economics.xlsx - instrumented admissions funnel by channel.

Route-by-route (Growth agent §11 note 1): Meta CTWA, Google Search, organic Instagram/content, referral/alumni,
paid Career Taster, webinar + regional pop-up. Target-driven like admissions_funnel.md §10: enrolment target x
route mix / route conversion = leads; leads x CPL = spend. Every calculation is a formula on Assumptions.
"""
import json
import os
import sys

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(__file__))
from xlhelpers import *  # noqa
import assumption_values as AV

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "economics", "funnel_economics.xlsx")
CASES = ("Conservative", "Base", "Aggressive")
FMTS = {"pct": FMT_PCT, "rs": FMT_RS, "int": FMT_INT, "1dp": FMT_1DP, "2dp": FMT_2DP, "gen": FMT_GEN}

FUNNEL_IDS = [k for k in AV.V]  # all shared values are relevant to the funnel


def write_assumptions(A, ids):
    groups = [("1. Canonical, baseline and GST", lambda k: k.startswith("CA-") or k in ("A-ECO-101", "A-ECO-102", "A-ECO-103", "A-ECO-23", "A-ECO-24")),
              ("2. Prices (Growth agent ladder L1-L5; BM realised fee)", lambda k: k in ("A-GRO-04", "A-GRO-05", "A-GRO-06", "A-BM-02", "A-GRO-07", "A-GRO-08", "A-GRO-09", "A-ECO-17")),
              ("3. Funnel volumes, mix, media costs and conversion (A-GRO master table)", lambda k: k.startswith("A-GRO") and k not in ("A-GRO-04", "A-GRO-05", "A-GRO-06", "A-GRO-07", "A-GRO-08", "A-GRO-09", "A-GRO-42", "A-GRO-45", "A-GRO-46", "A-GRO-47", "A-GRO-48", "A-GRO-49", "A-GRO-50")),
              ("4. Webinar / pop-up, follow-on and timing (A-BM, A-ECO)", lambda k: k.startswith("A-BM-06") or k in ("A-ECO-30", "A-ECO-31", "A-ECO-32", "A-ECO-33", "A-ECO-34", "A-ECO-35", "A-ECO-36", "A-ECO-37", "A-ECO-104", "A-ECO-105")),
              ("5. Delivery and growth-team costs (A-BM-04/16, A-GRO-42..50, A-ECO)", lambda k: k in ("A-BM-16", "A-BM-04a", "A-BM-04b", "A-BM-04c", "A-GRO-42", "A-GRO-45", "A-GRO-46", "A-GRO-47", "A-GRO-48", "A-GRO-49", "A-GRO-50") or k in ("A-ECO-40", "A-ECO-41", "A-ECO-42", "A-ECO-43", "A-ECO-44", "A-ECO-45", "A-ECO-46", "A-ECO-47", "A-ECO-48", "A-ECO-49"))]
    done = set()
    for gname, pred in groups:
        A.sec(gname)
        for k in ids:
            if k in done or not pred(k):
                continue
            d, c, b, g, unit, src, tag, f = AV.V[k]
            A.add(k, d, c, b, g, unit, src, tag, FMTS[f])
            done.add(k)
    rest = [k for k in ids if k not in done]
    if rest:
        A.sec("6. Other")
        for k in rest:
            d, c, b, g, unit, src, tag, f = AV.V[k]
            A.add(k, d, c, b, g, unit, src, tag, FMTS[f])


def funnel_block(wb, name, idx_formula, A, label):
    """One full funnel for a scenario. idx_formula: '=Assumptions!$C$4' or a constant 1/2/3."""
    ws = wb.create_sheet(name)
    title(ws, f"Instrumented funnel by channel - {label} (per month, steady state months 7-12)",
          "Target-driven: career enrolment target (A-GRO-01) x route mix / route conversion = leads; leads x cost = spend. Webinar/pop-up is the regional (S7) route, additive to the target.")
    ws["A3"] = "Scenario index"
    ws["B3"] = idx_formula
    ws["B3"].font = Font(bold=True, color="C00000")
    ws["C3"] = "=INDEX({\"Conservative\",\"Base\",\"Aggressive\"},B3)" if False else '=CHOOSE(B3,"Conservative","Base","Aggressive")'
    p = lambda aid: A.pick(aid, "$B$3")
    cols = ["Stage / metric", "Meta -> WhatsApp (CTWA)", "Google Search", "Organic IG / content / GBP", "Referral / alumni / walk-in",
            "Paid Career Taster", "Webinar + regional pop-up", "Total / blended", "Notes"]
    header(ws, 5, cols, widths=[46, 15, 15, 15, 15, 15, 15, 16, 56])
    R = {}
    RC = "BCDEF"  # four standard routes + taster live in B..F ; G webinar/pop-up ; H total
    r = 6

    def put(row, col, f, fmt, bold=False):
        c = ws.cell(row=row, column=col, value=f)
        c.number_format = fmt
        if bold:
            c.font = BOLD
            c.fill = KEY_FILL
        return c

    def label_row(row, text, notetxt="", bold=False):
        ws.cell(row=row, column=1, value=text).font = BOLD if bold else Font()
        ws.cell(row=row, column=9, value=notetxt).font = NOTE_FONT

    routes = [("A-GRO-18", "A-GRO-03a"), ("A-GRO-19", "A-GRO-03b"), ("A-GRO-20", "A-GRO-03c"), ("A-GRO-21", "A-GRO-03d")]

    section(ws, r, "A. Volumes", 9); r += 1
    label_row(r, "Career enrolment target per month (A-GRO-01)", "steady state, months 7-12")
    put(r, 8, "=" + p("A-GRO-01"), FMT_1DP, True); R["tgt"] = r; r += 1
    label_row(r, "Route mix (A-GRO-03)")
    for j, (_, mix) in enumerate(routes):
        put(r, 2 + j, "=" + p(mix), FMT_PCT)
    put(r, 6, "=" + p("A-GRO-03e"), FMT_PCT)
    put(r, 8, f"=SUM(B{r}:F{r})", FMT_PCT); R["mix"] = r; r += 1

    section(ws, r, "B. Stage conversion (scenario sets; base rates assume the <5-min SLA is met)", 9); r += 1
    stages = ["Lead -> Qualified lead", "QL -> Counselled / visit / demo", "Counselled -> Applied (token)", "Applied -> Enrolled (fee paid)"]
    for k, st in enumerate(stages):
        label_row(r, st, "Taster: ticket->attend / attend->counselled / counselled->enrolled ; Webinar/pop-up: see block F" if k == 0 else "")
        for j, (rid, _) in enumerate(routes):
            put(r, 2 + j, "=" + p(rid + "abcd"[k]), FMT_PCT)
        put(r, 6, "=" + p(["A-GRO-22a", "A-GRO-22b", "A-GRO-22c"][k]) if k < 3 else "=1", FMT_PCT)
        R[f"s{k}"] = r
        r += 1
    label_row(r, "Overall lead -> enrolled", "product of stages (Growth agent: CTWA 2.95/5.28/8.70%)", True)
    for j in range(5):
        col = "BCDEF"[j]
        put(r, 2 + j, f"={col}{R['s0']}*{col}{R['s1']}*{col}{R['s2']}*{col}{R['s3']}", FMT_PCT2, True)
    R["ov"] = r; r += 1

    section(ws, r, "C. Funnel counts per month", 9); r += 1
    label_row(r, "Enrolled (target x mix)")
    for j in range(5):
        col = "BCDEF"[j]
        put(r, 2 + j, f"=$H${R['tgt']}*{col}{R['mix']}", FMT_1DP)
    R["enr"] = r; r += 1
    label_row(r, "Leads (Taster: tickets sold)", "enrolled / overall rate")
    for j in range(5):
        col = "BCDEF"[j]
        put(r, 2 + j, f"=IFERROR({col}{R['enr']}/{col}{R['ov']},0)", FMT_1DP)
    R["lead"] = r; r += 1
    for k, nm in enumerate(["Qualified leads (Taster: attended)", "Counselled / visited", "Applied (token paid)"]):
        label_row(r, nm)
        for j in range(5):
            col = "BCDEF"[j]
            prev = R["lead"] if k == 0 else R[f"c{k - 1}"]
            put(r, 2 + j, f"={col}{prev}*{col}{R['s' + str(k)]}", FMT_1DP)
        R[f"c{k}"] = r; r += 1
    label_row(r, "Enrolled (check = target x mix)")
    for j in range(5):
        col = "BCDEF"[j]
        put(r, 2 + j, f"={col}{R['c2']}*{col}{R['s3']}", FMT_1DP)
    R["enrChk"] = r; r += 1

    section(ws, r, "D. Cost per lead and channel spend (₹ per month)", 9); r += 1
    label_row(r, "Cost per lead / ticket", "CTWA blends tricity and regional CPL by out-of-town share; organic = content cost / leads; referral = reward per enrolment")
    put(r, 2, f"=(1-{p('A-GRO-37')})*{p('A-GRO-10')}+{p('A-GRO-37')}*{p('A-GRO-11')}", FMT_RS)
    put(r, 3, "=" + p("A-GRO-12"), FMT_RS)
    put(r, 6, "=" + p("A-GRO-15"), FMT_RS)
    R["cpl"] = r; r += 1
    label_row(r, "Media / reward / content spend", "Organic: shooter/editor (A-GRO-47) + SEO (A-GRO-50); fixed team cost, shown for CAC visibility")
    put(r, 2, f"=B{R['lead']}*B{R['cpl']}", FMT_RS)
    put(r, 3, f"=C{R['lead']}*C{R['cpl']}", FMT_RS)
    put(r, 4, "=20000+15000", FMT_RS)  # placeholder replaced below
    put(r, 5, f"=E{R['enr']}*{p('A-GRO-34')}", FMT_RS)
    put(r, 6, f"=F{R['lead']}*(F{R['cpl']}+{p('A-ECO-44')})-F{R['lead']}*{p('A-GRO-04')}/(1+{p('A-ECO-23')})", FMT_RS)
    R["spend"] = r; r += 1
    ws.cell(row=R["spend"], column=4, value="=CONTENT_COST")  # patched after Assumptions exist
    label_row(r, "  of which taster ticket revenue offset (ex-GST)", "Growth §11 note 3")
    put(r, 6, f"=-F{R['lead']}*{p('A-GRO-04')}/(1+{p('A-ECO-23')})", FMT_RS)
    R["tick"] = r; r += 1

    section(ws, r, "E. Unit economics per enrolled career student (₹, ex-GST)", 9); r += 1
    gst = f"(1+{p('A-ECO-23')}*{p('A-ECO-24')})"
    l3fee = f"IF({p('A-ECO-104')}<=7,{p('A-BM-02')},{p('CA-02')})"
    label_row(r, "Steady-state L3 realised fee (incl. GST if quoted inclusive)", "A-BM-02 after price step (A-ECO-104 <= month 7), else CA-02")
    put(r, 8, "=" + l3fee, FMT_RS); R["l3"] = r; r += 1
    label_row(r, "First-course revenue per career enrolment (ex-GST)", "L3 share x L3 + L2 share x L2; taster route nets the ₹999 credit")
    for j in range(6):
        col = "BCDEFG"[j]
        credit = f"-{p('A-GRO-04')}" if col == "F" else ""
        put(r, 2 + j, f"=({p('A-GRO-02')}*$H${R['l3']}+(1-{p('A-GRO-02')})*{p('A-GRO-05')}{credit})/{gst}", FMT_RS)
    R["rev"] = r; r += 1
    trainer_std = f"({p('A-BM-04b')}+{p('A-BM-04c')})/{p('A-ECO-46')}/{p('A-ECO-47')}"
    label_row(r, "Trainer standard cost per seat-month (trainer + assistant / slots / fill)", "A-BM-04b/c, A-ECO-46/47")
    put(r, 8, "=" + trainer_std, FMT_RS); R["tstd"] = r; r += 1
    var = (f"{p('A-GRO-02')}*({p('A-BM-16')}+IF({p('A-ECO-104')}<=7,0,{p('A-ECO-103')})+{p('A-ECO-48')}*$H${R['tstd']}+$H${R['l3']}/{gst}*{p('A-GRO-35')}*{p('A-GRO-35b')})"
           f"+(1-{p('A-GRO-02')})*({p('A-ECO-40')}+$H${R['tstd']})")
    label_row(r, "Variable delivery cost per career enrolment", "consumables + kit (pre-step) + trainer standard cost + EMI subvention; gateway below")
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f"={var}+{col}{R['rev']}*{p('A-GRO-41')}", FMT_RS)
    R["var"] = r; r += 1
    label_row(r, "First-course contribution before CAC", "", True)
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f"={col}{R['rev']}-{col}{R['var']}", FMT_RS, True)
    R["cm"] = r; r += 1
    label_row(r, "Contribution margin % (before CAC)")
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f"=IFERROR({col}{R['cm']}/{col}{R['rev']},0)", FMT_PCT)
    R["cmp"] = r; r += 1
    fol = (f"(1-{p('A-GRO-02')})*{p('A-ECO-36')}*(({p('A-BM-02')}-{p('A-GRO-05')})/{gst}-{p('A-BM-16')}-2*$H${R['tstd']})"
           f"+{p('A-ECO-35')}*({p('A-GRO-07')}/{gst}-{p('A-ECO-42')})"
           f"+{p('A-GRO-31')}*{p('A-GRO-36')}*({p('A-GRO-09')}/{gst}-{p('A-ECO-45')})"
           f"+{p('A-GRO-02')}*IF({p('A-ECO-104')}<=7,{p('A-ECO-37')},0)")
    label_row(r, "Follow-on contribution within 24 months (LTV add-on)", "L2->L3 upgrade + masterclasses + Pro Circle renewal + kit margin")
    put(r, 8, "=" + fol, FMT_RS); R["fol"] = r; r += 1
    label_row(r, "LTV: 24-month contribution per career student (before CAC)", "", True)
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f"={col}{R['cm']}+$H${R['fol']}", FMT_RS, True)
    R["ltv"] = r; r += 1

    section(ws, r, "F. Webinar + regional pop-up route (S7; additive to target)", 9); r += 1
    label_row(r, "Online parents' sessions per month x registrants", "A-ECO-31 x A-ECO-32")
    put(r, 7, f"={p('A-ECO-31')}*{p('A-ECO-32')}", FMT_1DP); R["wreg"] = r; r += 1
    label_row(r, "Webinar enrolments = registrants x attend x counselling booked x CTWA counselled->enrolled", "A-GRO-26a/b; A-GRO-18c/d")
    put(r, 7, f"=G{R['wreg']}*{p('A-GRO-26a')}*{p('A-GRO-26b')}*{p('A-GRO-18c')}*{p('A-GRO-18d')}", FMT_2DP); R["wenr"] = r; r += 1
    label_row(r, "Pop-up events per month (A-ECO-33) x paid attendees", "A-BM-06a")
    put(r, 7, f"={p('A-ECO-33')}*{p('A-BM-06a')}", FMT_1DP); R["patt"] = r; r += 1
    label_row(r, "Pop-up enrolments = attendees x attendee->enrol (A-BM-06d)")
    put(r, 7, f"=G{R['patt']}*{p('A-BM-06d')}", FMT_2DP); R["penr"] = r; r += 1
    label_row(r, "Webinar promo + net pop-up cost (event cost - ticket revenue ex-GST)")
    put(r, 7, f"=G{R['wreg']}*{p('A-ECO-30')}+{p('A-ECO-33')}*({p('A-BM-06c')}-{p('A-BM-06a')}*{p('A-BM-06b')}/(1+{p('A-ECO-23')}))", FMT_RS); R["wcost"] = r; r += 1
    # roll webinar/pop-up into the main rows
    ws.cell(row=R["enr"], column=7, value=f"=G{R['wenr']}+G{R['penr']}").number_format = FMT_1DP
    ws.cell(row=R["enrChk"], column=7, value=f"=G{R['enr']}").number_format = FMT_1DP
    ws.cell(row=R["lead"], column=7, value=f"=G{R['wreg']}+G{R['patt']}").number_format = FMT_1DP
    ws.cell(row=R["spend"], column=7, value=f"=G{R['wcost']}").number_format = FMT_RS
    ws.cell(row=R["ov"], column=7, value=f"=IFERROR(G{R['enr']}/G{R['lead']},0)").number_format = FMT_PCT2

    section(ws, r, "G. CAC, payback and totals", 9); r += 1
    label_row(r, "CAC per enrolment (channel spend / enrolled)", "", True)
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f"=IFERROR({col}{R['spend']}/{col}{R['enr']},0)", FMT_RS, True)
    R["cac"] = r; r += 1
    label_row(r, "CAC as % of first-course revenue")
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f"=IFERROR({col}{R['cac']}/{col}{R['rev']},0)", FMT_PCT)
    R["cacp"] = r; r += 1
    label_row(r, "LTV / CAC (24-month contribution)")
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f'=IF({col}{R["cac"]}<=0,"n/m (self-funding)",{col}{R["ltv"]}/{col}{R["cac"]})', FMT_2DP)
    R["ltvcac"] = r; r += 1
    label_row(r, "CAC payback (months) = CAC / (first-course contribution / payment months)", "BM M3 definition; A-ECO-34 payment months")
    for j in range(6):
        col = "BCDEFG"[j]
        put(r, 2 + j, f"=IF({col}{R['cac']}<=0,0,{col}{R['cac']}/({col}{R['cm']}/{p('A-ECO-34')}))", FMT_1DP)
    R["pb"] = r; r += 1

    # totals column H
    tot = lambda key, fmt: put(R[key], 8, f"=SUM(B{R[key]}:G{R[key]})", fmt, True)
    for key in ("enr", "lead", "c0", "c1", "c2", "enrChk", "spend"):
        tot(key, FMT_1DP if key != "spend" else FMT_RS)
    put(R["ov"], 8, f"=H{R['enr']}/H{R['lead']}", FMT_PCT2, True)
    put(R["rev"], 8, f"=SUMPRODUCT(B{R['rev']}:G{R['rev']},B{R['enr']}:G{R['enr']})/H{R['enr']}", FMT_RS, True)
    put(R["var"], 8, f"=SUMPRODUCT(B{R['var']}:G{R['var']},B{R['enr']}:G{R['enr']})/H{R['enr']}", FMT_RS)
    put(R["cm"], 8, f"=H{R['rev']}-H{R['var']}", FMT_RS, True)
    put(R["cmp"], 8, f"=H{R['cm']}/H{R['rev']}", FMT_PCT)
    put(R["ltv"], 8, f"=H{R['cm']}+H{R['fol']}", FMT_RS, True)
    put(R["cac"], 8, f"=H{R['spend']}/H{R['enr']}", FMT_RS, True)
    put(R["cacp"], 8, f"=H{R['cac']}/H{R['rev']}", FMT_PCT)
    put(R["ltvcac"], 8, f"=H{R['ltv']}/H{R['cac']}", FMT_2DP, True)
    put(R["pb"], 8, f"=H{R['cac']}/(H{R['cm']}/{p('A-ECO-34')})", FMT_1DP, True)

    label_row(r, "Paid-only CAC (CTWA + Search + Taster net + webinar/pop-up)", "excludes organic content and referral rewards", True)
    put(r, 8, f"=(B{R['spend']}+C{R['spend']}+F{R['spend']}+G{R['spend']})/(B{R['enr']}+C{R['enr']}+F{R['enr']}+G{R['enr']})", FMT_RS, True); R["paidcac"] = r; r += 1
    label_row(r, "Variable acquisition cost per target career enrolment (for business_forecast)", "(CTWA + Search + Taster net + referral rewards) / target enrolments; organic content is a fixed team cost there", True)
    put(r, 8, f"=(B{R['spend']}+C{R['spend']}+E{R['spend']}+F{R['spend']})/H{R['tgt']}", FMT_RS, True); R["fcacq"] = r; r += 1
    label_row(r, "Paid-route media CAC: (CTWA + Search + Taster net) / their enrolments (for business_forecast)", "Growth G4 'media CAC per paid career enrolment'", True)
    put(r, 8, f"=(B{R['spend']}+C{R['spend']}+F{R['spend']})/(B{R['enr']}+C{R['enr']}+F{R['enr']})", FMT_RS, True); R["prcac"] = r; r += 1
    label_row(r, "Counselling conversations per month (career leads x 2.5 + other leads x 0.5)", "A-GRO §10.4 method")
    put(r, 8, f"=H{R['lead']}*2.5+H{{OTHER}}*0.5", FMT_INT); R["conv"] = r; r += 1
    label_row(r, "Counsellor FTE needed (A-GRO-38 x 24 days)")
    put(r, 8, f"=H{R['conv']}/({p('A-GRO-38')}*24)", FMT_2DP, True); R["fte"] = r; r += 1
    label_row(r, "Fully loaded CAC incl. counsellor, performance freelancer and martech", "(channel spend + A-GRO-45 x FTE + A-GRO-48 + A-GRO-42) / enrolments; Growth G4 definition")
    put(r, 8, f"=(H{R['spend']}+MAX(1,ROUNDUP(H{R['fte']},0))*{p('A-GRO-45')}+{p('A-GRO-48')}+{p('A-GRO-42')})/H{R['enr']}", FMT_RS, True); R["flcac"] = r; r += 1
    label_row(r, "Media spend per month on career routes (CTWA + Search + Taster + pop-up/webinar)")
    put(r, 8, f"=B{R['spend']}+C{R['spend']}+F{R['spend']}+G{R['spend']}", FMT_RS, True); R["media"] = r; r += 1
    label_row(r, "Check: organic leads derived vs A-GRO-43 ; referral leads vs A-GRO-44", "Growth agent check values")
    put(r, 4, f"=D{R['lead']}-{p('A-GRO-43')}", FMT_1DP)
    put(r, 5, f"=E{R['lead']}-{p('A-GRO-44')}", FMT_1DP)
    R["chk"] = r; r += 2

    section(ws, r, "H. Other products: masterclass (L4) and self-makeup (L1-S5) acquisition", 9); r += 1
    label_row(r, "Seats per month (A-GRO §10.3: L4 24 base; L1-S5 15 base) - scaled with target", "L4 = 12 x target/5 ; L1-S5 = 1.5 x target")
    put(r, 2, f"=2.4*$H${R['tgt']}", FMT_1DP)
    put(r, 3, f"=1.5*$H${R['tgt']}", FMT_1DP)
    ws.cell(row=r, column=2).comment = None
    R["seats"] = r; r += 1
    label_row(r, "Lead -> purchase (A-GRO-27 / A-GRO-28)")
    put(r, 2, "=" + p("A-GRO-27"), FMT_PCT); put(r, 3, "=" + p("A-GRO-28"), FMT_PCT); R["oc"] = r; r += 1
    label_row(r, "Leads needed")
    put(r, 2, f"=B{R['seats']}/B{R['oc']}", FMT_1DP); put(r, 3, f"=C{R['seats']}/C{R['oc']}", FMT_1DP)
    put(r, 8, f"=B{r}+C{r}", FMT_1DP); R["olead"] = r; r += 1
    label_row(r, "Paid media (paid share x CPL x leads)", "A-GRO-18b, A-GRO-16/17")
    put(r, 2, f"=B{R['olead']}*{p('A-GRO-18b-MC')}*{p('A-GRO-16')}", FMT_RS); put(r, 3, f"=C{R['olead']}*{p('A-GRO-18b-SM')}*{p('A-GRO-17')}", FMT_RS)
    R["omedia"] = r; r += 1
    label_row(r, "Media per seat sold (for business_forecast)", "", True)
    put(r, 2, f"=B{R['omedia']}/B{R['seats']}", FMT_RS, True); put(r, 3, f"=C{R['omedia']}/C{R['seats']}", FMT_RS, True)
    R["oper"] = r; r += 1
    label_row(r, "Columns: B = L4 masterclass, C = L1-S5 self-makeup")
    # patch counsellor-load formula with the other-leads row
    ws.cell(row=R["conv"], column=8, value=f"=H{R['lead']}*2.5+H{R['olead']}*0.5")
    return ws, R


def build():
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    A = Assumptions(wb.create_sheet(), case_labels=CASES, selector_default="Base",
                    intro="Blue = input. Columns are scenario values: Conservative = adverse rates + smaller plan; Aggressive = favourable rates + larger plan. A-GRO = Growth agent master table; A-BM = Business Model agent; A-ECO = Economics agent.")
    write_assumptions(A, FUNNEL_IDS)
    blocks = {}
    for nm, idx, lab in [("Funnel_Selected", "=Assumptions!$C$4", "scenario chosen on Assumptions!C3"),
                         ("Funnel_Conservative", 1, "Conservative"), ("Funnel_Base", 2, "Base"), ("Funnel_Aggressive", 3, "Aggressive")]:
        ws, R = funnel_block(wb, nm, idx, A, lab)
        ws.cell(row=R["spend"], column=4, value=f"={A.pick('A-GRO-47', '$B$3')}+{A.pick('A-GRO-50', '$B$3')}").number_format = FMT_RS
        blocks[nm] = R

    # ---------------- Speed_to_Lead
    R = blocks["Funnel_Base"]
    ws = wb.create_sheet("Speed_to_Lead")
    title(ws, "Effect of speed-to-lead discipline (H7) - Base scenario, per month",
          "Same enrolment target with and without the <5-min SLA (A-GRO-24 multiplies Lead->QL when the SLA is missed), and the alternative view: enrolments lost at the same media budget.")
    header(ws, 4, ["Metric", "Meta CTWA", "Google Search", "Organic", "Total (3 lead routes)", "Notes"], widths=[52, 15, 15, 15, 18, 60])
    FB = "Funnel_Base!"
    sla = A.r("A-GRO-24", "base")
    rows = [
        ("Enrolments target (routes as Funnel_Base)", lambda c: f"={FB}{c}{R['enr']}", FMT_1DP),
        ("Leads needed - SLA met", lambda c: f"={FB}{c}{R['lead']}", FMT_1DP),
        ("Leads needed - SLA missed (Lead->QL x A-GRO-24)", lambda c: f"={FB}{c}{R['lead']}/{sla}", FMT_1DP),
        ("Spend - SLA met (₹)", lambda c: f"={FB}{c}{R['spend']}", FMT_RS),
        ("Spend - SLA missed (₹)", lambda c: f"={FB}{c}{R['spend']}/{sla}", FMT_RS),
        ("CAC - SLA met (₹)", lambda c: f"={FB}{c}{R['cac']}", FMT_RS),
        ("CAC - SLA missed (₹)", lambda c: f"={FB}{c}{R['cac']}/{sla}", FMT_RS),
        ("Extra spend per month to hit the same target (₹)", lambda c: f"={FB}{c}{R['spend']}*(1/{sla}-1)", FMT_RS),
        ("Alternative: enrolments at the SAME budget if SLA missed", lambda c: f"={FB}{c}{R['enr']}*{sla}", FMT_1DP),
        ("Enrolments lost per month at the same budget", lambda c: f"={FB}{c}{R['enr']}*(1-{sla})", FMT_1DP),
        ("Contribution lost per month (lost enrolments x contribution)", lambda c: f"={FB}{c}{R['enr']}*(1-{sla})*{FB}{c}{R['cm']}", FMT_RS),
    ]
    for i, (lab, fn, fmt) in enumerate(rows):
        rr = 5 + i
        ws.cell(row=rr, column=1, value=lab)
        for j, c in enumerate("BCD"):
            ws.cell(row=rr, column=2 + j, value=fn(c)).number_format = fmt
        ws.cell(row=rr, column=5, value=f"=SUM(B{rr}:D{rr})" if "CAC" not in lab else f"=IFERROR(SUMPRODUCT(B{rr}:D{rr},B5:D5)/SUM(B5:D5),0)").number_format = fmt
        ws.cell(row=rr, column=5).font = BOLD
    ws.cell(row=17, column=1, value="Annual contribution lost at the same budget (x12)").font = BOLD
    ws.cell(row=17, column=5, value="=E15*12").number_format = FMT_RS
    ws.cell(row=18, column=1, value="Founder-touch: counselled->applied falls by 1/A-GRO-25 when fully team-led; enrolments lost at same leads (all routes)").font = BOLD
    ws.cell(row=18, column=5, value=f"=Funnel_Base!H{R['enr']}*(1-1/{A.r('A-GRO-25', 'base')})").number_format = FMT_1DP
    note(ws, 20, "Reading: missing the SLA costs ~1/0.67 = 1.5x the media for the same enrolments - more than the gap between base and optimistic CPL. Process discipline before budget (H7).")

    # ---------------- Popup_Events (BM M5)
    wp = wb.create_sheet("Popup_Events")
    title(wp, "Regional pop-up discovery workshop economics (per event) - BM spec M5")
    header(wp, 4, ["Metric", "Conservative", "Base", "Aggressive", "Notes"], widths=[52, 15, 15, 15, 60])
    pr = lambda aid, j: A.r(aid, j + 1)
    prow = [
        ("Paid attendees", lambda j: f"={pr('A-BM-06a', j)}", FMT_INT, "A-BM-06a"),
        ("Ticket price (incl. GST)", lambda j: f"={pr('A-BM-06b', j)}", FMT_RS, "A-BM-06b"),
        ("Ticket revenue ex-GST", lambda j: f"={'BCD'[j]}5*{'BCD'[j]}6/(1+{pr('A-ECO-23', j)})", FMT_RS, ""),
        ("Event cost", lambda j: f"={pr('A-BM-06c', j)}", FMT_RS, "A-BM-06c"),
        ("Event net result (before enrolments)", lambda j: f"={'BCD'[j]}7-{'BCD'[j]}8", FMT_RS, ""),
        ("Breakeven paid attendance", lambda j: f"={'BCD'[j]}8/({'BCD'[j]}6/(1+{pr('A-ECO-23', j)}))", FMT_1DP, "BM: ~36 at ₹1,499"),
        ("Enrolments per event (attendee->enrol A-BM-06d)", lambda j: f"={'BCD'[j]}5*{pr('A-BM-06d', j)}", FMT_2DP, ""),
        ("CAC per enrolment (net event cost / enrolments)", lambda j: f"=IF({'BCD'[j]}9>=0,0,-{'BCD'[j]}9/{'BCD'[j]}11)", FMT_RS, "0 = event self-funding"),
        ("Attendee->enrol needed for CAC <= ₹8,000", lambda j: f"=MAX(0,-{'BCD'[j]}9)/(8000*{'BCD'[j]}5)", FMT_PCT, "BM M5 threshold"),
        ("Contribution from enrolments (x Funnel career contribution)", lambda j: f"={'BCD'[j]}11*{['Funnel_Conservative', 'Funnel_Base', 'Funnel_Aggressive'][j]}!G{R['cm']}", FMT_RS, ""),
        ("Event total contribution incl. enrolments", lambda j: f"={'BCD'[j]}9+{'BCD'[j]}14", FMT_RS, ""),
    ]
    for i, (lab, fn, fmt, nt) in enumerate(prow):
        rr = 5 + i
        wp.cell(row=rr, column=1, value=lab)
        for j in range(3):
            wp.cell(row=rr, column=2 + j, value=fn(j)).number_format = fmt
        wp.cell(row=rr, column=5, value=nt).font = NOTE_FONT

    # ---------------- Summary
    wsu = wb.create_sheet("Summary", 1)
    title(wsu, "Funnel economics - scenario summary (per month at steady state unless stated)",
          "Read with admissions_funnel.md §10 (Growth agent). Outputs in the yellow rows feed business_forecast.xlsx (pasted by run_all.py).")
    header(wsu, 4, ["Metric", "Conservative", "Base", "Aggressive", "Selected", "Source row"], widths=[60, 16, 16, 16, 16, 30])
    items = [("Career enrolment target / month", "tgt", "H", FMT_1DP), ("Career + webinar/pop-up enrolments / month", "enr", "H", FMT_1DP),
             ("Leads / month (career routes + webinar/pop-up)", "lead", "H", FMT_INT), ("Overall lead -> enrolled", "ov", "H", FMT_PCT2),
             ("Media spend / month, career routes (₹)", "media", "H", FMT_RS), ("Channel spend incl. content + rewards (₹)", "spend", "H", FMT_RS),
             ("Meta CTWA CAC (₹)", "cac", "B", FMT_RS), ("Google Search CAC (₹)", "cac", "C", FMT_RS), ("Organic CAC (content cost) (₹)", "cac", "D", FMT_RS),
             ("Referral CAC (reward) (₹)", "cac", "E", FMT_RS), ("Taster CAC net of tickets (₹)", "cac", "F", FMT_RS), ("Webinar + pop-up CAC (₹)", "cac", "G", FMT_RS),
             ("Blended CAC, all routes (₹)", "cac", "H", FMT_RS), ("Paid-only CAC (₹)", "paidcac", "H", FMT_RS), ("Fully loaded CAC incl. counsellor (₹)", "flcac", "H", FMT_RS),
             ("Blended CAC as % of first-course revenue", "cacp", "H", FMT_PCT), ("First-course revenue per career student, ex-GST (₹)", "rev", "H", FMT_RS),
             ("First-course contribution before CAC (₹)", "cm", "H", FMT_RS), ("Follow-on 24-month contribution (₹)", "fol", "H", FMT_RS),
             ("LTV (24-month contribution) (₹)", "ltv", "H", FMT_RS), ("LTV / blended CAC", "ltvcac", "H", FMT_2DP), ("CAC payback (months)", "pb", "H", FMT_1DP),
             ("Counsellor FTE needed", "fte", "H", FMT_2DP),
             ("FEED: variable acquisition cost per target career enrolment (₹)", "fcacq", "H", FMT_RS),
             ("FEED: paid-route media CAC, CTWA + Search + Taster (₹)", "prcac", "H", FMT_RS),
             ("Funnel-required career media / month at target (₹)", "media", "H", FMT_RS),
             ("FEED: masterclass media per seat (₹)", "oper", "B", FMT_RS), ("FEED: self-makeup media per seat (₹)", "oper", "C", FMT_RS)]
    SR = {}
    for i, (lab, key, col, fmt) in enumerate(items):
        rr = 5 + i
        wsu.cell(row=rr, column=1, value=lab).font = BOLD if lab.startswith("FEED") else Font()
        for j, sh in enumerate(["Funnel_Conservative", "Funnel_Base", "Funnel_Aggressive", "Funnel_Selected"]):
            c = wsu.cell(row=rr, column=2 + j, value=f"={sh}!{col}{R[key]}")
            c.number_format = fmt
            if lab.startswith("FEED"):
                c.fill = KEY_FILL
        wsu.cell(row=rr, column=6, value=f"{col}{R[key]}").font = NOTE_FONT
        SR[lab] = rr
    note(wsu, 5 + len(items) + 1, "CAC payback: BM gate G1 requires CAC payback <= 3 months and CAC <= 12% of realised fee.")

    # ---------------- README
    title(readme, "PNA funnel economics - instrumented admissions funnel by channel",
          "Economics & Market Sizing agent, 2026-10-16. Build script: economics/build_scripts/build_funnel_economics.py")
    lines = [("Alignment", "Uses the Growth agent's master assumption table (strategy/admissions_funnel.md §9, A-GRO-01..50) with its IDs and values, mapped Low/Base/High -> scenario columns per its convention (cost 'High' = Conservative). Adds A-BM (Business Model agent) pop-up inputs and A-ECO delivery-cost inputs."),
             ("Method", "Target-driven per month at steady state (months 7-12): A-GRO-01 x route mix (A-GRO-03) / route conversion (A-GRO-18..22) = leads; leads x CPL = spend. Webinar + regional pop-up route (S7) is additive. Unit economics per career student: first-course revenue (ex-GST) - variable delivery cost = contribution; + 24-month follow-on = LTV."),
             ("Sheets", "Summary (Cons/Base/Aggr side by side) ; Assumptions ; Funnel_Selected (driven by Assumptions!C3) ; Funnel_Conservative / _Base / _Aggressive ; Speed_to_Lead (H7) ; Popup_Events (BM M5)."),
             ("Feeds", "Summary 'FEED' rows -> business_forecast.xlsx Assumptions A-ECO-FUN-01..03 (pasted by run_all.py)."),
             ("GST", "Fees assumed quoted GST-inclusive (A-ECO-24 = 1): revenue = fee / 1.18. 18% GST applies to private non-recognised training (A-ECO-23)."),
             ("Limitations", "All conversion and CPL inputs are hypotheses from vendor/US benchmarks (A-CUS/A-GRO); replace with PNA CRM data after 6-8 weeks of instrumentation. Organic and referral volumes are derived from the target, not observed.")]
    for i, (k, v) in enumerate(lines):
        readme.cell(row=4 + i, column=1, value=k).font = BOLD
        c = readme.cell(row=4 + i, column=2, value=v)
        c.alignment = WRAP
        readme.row_dimensions[4 + i].height = 48
    setw(readme, {"A": 16, "B": 120})
    wb.save(OUT)
    return OUT, dict(R=R, SR=SR)


if __name__ == "__main__":
    out, refs = build()
    json.dump(refs, open(os.path.join(os.path.dirname(__file__), "_funnel_refs.json"), "w"), indent=1)
    print(out)
